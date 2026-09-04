"""objective_analyze 插件测试。"""

from __future__ import annotations

from datetime import date as _date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from grade_analyzer.config import load_config
from plugins.objective_analyze import analyze

FIXTURES = Path(__file__).parent / "fixtures" / "objective_analyze"
INPUT_DIR = FIXTURES / "限时练三(地理)客观题得分明细"
EXAM_DIR = "高一下地理限时练三"
SUM_NAME = "高一下地理限时练三_客观题得分汇总（全部班级）.xlsx"
DEV_NAME = "高一下地理限时练三_客观题得分率距平（柯）.xlsx"


def _sum_path(out_dir: Path) -> Path:
    return out_dir / EXAM_DIR / SUM_NAME


def _sum_named(out_dir: Path, label: str) -> Path:
    return out_dir / EXAM_DIR / f"高一下地理限时练三_客观题得分汇总（{label}）.xlsx"


def _dev_path(out_dir: Path) -> Path:
    return out_dir / EXAM_DIR / DEV_NAME


def _dev_named(out_dir: Path, label: str) -> Path:
    return out_dir / EXAM_DIR / f"高一下地理限时练三_客观题得分率距平（{label}）.xlsx"


def _write_configs(
    tmp_path: Path,
    with_exam: bool = False,
    exam_date: str = "",
    output_dir_cfg: str = "",
    subdir_by_exam: bool = True,
    summary_groups: str = "[]",
    deviation_groups: str = "[teacher]",
    baseline: str = "",
    level_classes: str = (
        "高一1班: {level: A, course: 物化地}\n"
        "高一2班: {level: B, course: 史地政}\n"
    ),
) -> tuple[Path, Path, Path]:
    """写临时全局配置、学科配置与插件配置。"""
    subjects_dir = tmp_path / "subjects"
    subjects_dir.mkdir()
    classes_dir = tmp_path / "classes"
    classes_dir.mkdir()
    exams_dir = tmp_path / "exams"
    exams_dir.mkdir()
    (subjects_dir / "高一第二学期_地理.yaml").write_text(
        "subject: 地理\n"
        "teacher_count: 1\n"
        "teacher_names:\n"
        "  A: 柯\n"
        "class_teachers:\n"
        "  高一1班: A\n"
        "  高一2班: A\n",
        encoding="utf-8",
    )
    (classes_dir / "高一第二学期.yaml").write_text(level_classes, encoding="utf-8")
    if with_exam:
        (exams_dir / "高一第二学期").mkdir()
        (exams_dir / "高一第二学期" / "限时练三.yaml").write_text(
            "subject: 地理\n"
            "semester: 高一第二学期\n"
            "date: '2026-05-20'\n"
            "folder: data/input\n"
            "file: 不存在.xlsx\n"
            "name: 高一下地理限时练三\n"
            "short_name: 限时练三\n",
            encoding="utf-8",
        )
    global_cfg = tmp_path / "config.yaml"
    global_cfg.write_text(
        "subjects: [语文, 数学, 外语, 物理, 化学, 生物, 政治, 历史, 地理, 技术]\n"
        f"current_semester: 高一第二学期\n"
        f"subjects_dir: '{subjects_dir}'\n"
        f"classes_dir: '{classes_dir}'\n"
        f"exams_dir: '{exams_dir}'\n",
        encoding="utf-8",
    )
    plugin_cfg = tmp_path / "plugin.yaml"
    plugin_cfg.write_text(
        "input_dir: ''\n"
        f"output_dir: '{output_dir_cfg}'\n"
        "semester: ''\n"
        "subject: ''\n"
        f"exam_date: '{exam_date}'\n"
        "low_score_flag: 0.6\n"
        "max_date_input_errors: 2\n"
        f"output_subdir_by_exam: {str(subdir_by_exam).lower()}\n"
        f"summary_groups: {summary_groups}\n"
        f"deviation_groups: {deviation_groups}\n"
        f"baseline: '{baseline}'\n",
        encoding="utf-8",
    )
    return global_cfg, plugin_cfg, tmp_path / "out"


def test_trans_num():
    assert analyze.trans_num("47.92%") == 0.4792
    assert analyze.trans_num("0.96") == 0.96
    assert analyze.trans_num("2.0") == 2.0
    assert analyze.trans_num("D") == "D"
    assert analyze.trans_num("") is None


def test_load_plugin_config_default_auto():
    cfg = analyze._load_plugin_config(None)
    assert "summary_groups" in cfg
    assert "deviation_groups" in cfg
    assert "baseline" in cfg


def test_load_plugin_config_bl_alias(tmp_path):
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text("baseline: ''\nbl: '柯'\n", encoding="utf-8")
    raw = analyze._load_plugin_config(str(cfg))
    assert raw["bl"] == "柯"


def test_load_plugin_config_explicit_missing_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        analyze._load_plugin_config(str(tmp_path / "nope.yaml"))


def test_split_folder_name():
    assert analyze._split_folder_name("限时练三(地理)客观题得分明细") == ("限时练三", "地理")
    assert analyze._split_folder_name("限时练三（地理）客观题得分明细") == ("限时练三", "地理")
    assert analyze._split_folder_name("限时练三客观题得分明细") == ("限时练三", None)


def test_parse_date():
    assert analyze._parse_date("2026-5-20") == "2026-05-20"
    assert analyze._parse_date("2026/5/20") == "2026-05-20"
    assert analyze._parse_date("2026.5.20") == "2026-05-20"
    assert analyze._parse_date("2026年5月20日") == "2026-05-20"
    with pytest.raises(ValueError, match="格式不正确"):
        analyze._parse_date("abc")
    with pytest.raises(ValueError, match="不存在"):
        analyze._parse_date("2026-13-20")


def test_derive_exam_info(tmp_path):
    global_cfg, _plugin_cfg, _out = _write_configs(tmp_path)
    cfg = load_config(str(global_cfg))
    exam_name, subject = analyze.derive_exam_info(
        "限时练三(地理)客观题得分明细", "高一第二学期", cfg
    )
    assert subject == "地理"
    assert exam_name == "高一下地理限时练三"
    exam_name2, subject2 = analyze.derive_exam_info(
        "限时练三客观题得分明细", "高一第二学期", cfg
    )
    assert subject2 is None
    assert exam_name2 == "高一下限时练三"


def test_teachers_from_config(tmp_path):
    global_cfg, _plugin_cfg, _out = _write_configs(tmp_path)
    cfg = load_config(str(global_cfg))
    assert analyze.teachers_from_config(cfg, "高一第二学期", "地理") == {
        "柯": ["1班", "2班"]
    }


def test_build_groups(tmp_path):
    global_cfg, _plugin_cfg, _out = _write_configs(tmp_path)
    cfg = load_config(str(global_cfg))
    groups = analyze.build_groups(
        cfg, "高一第二学期", "地理", ["1班", "2班"], ["teacher", "level"]
    )
    assert [(g.group_type, g.name, g.classes) for g in groups["teacher"]] == [
        ("teacher", "柯", ["1班", "2班"])
    ]
    assert [(g.group_type, g.name, g.classes) for g in groups["level"]] == [
        ("level", "A", ["1班"]),
        ("level", "B", ["2班"]),
    ]


def test_build_groups_unassigned(tmp_path):
    global_cfg, _plugin_cfg, _out = _write_configs(
        tmp_path, level_classes="高一1班: {level: A, course: 物化地}\n"
    )
    cfg = load_config(str(global_cfg))
    groups = analyze.build_groups(cfg, "高一第二学期", "地理", ["1班", "2班"], ["level"])
    assert [(g.name, g.classes) for g in groups["level"]] == [
        ("A", ["1班"]),
        ("未分层", ["2班"]),
    ]


def test_run_generates_outputs(tmp_path):
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path, with_exam=True)
    result = analyze.run(
        config_path=str(global_cfg),
        plugin_config=str(plugin_cfg),
        input_dir=str(INPUT_DIR),
        output_dir=str(out_dir),
    )
    assert result == 0

    sum_file = _sum_path(out_dir)
    dev_file = _dev_path(out_dir)
    assert sum_file.is_file()
    assert dev_file.is_file()

    wb = load_workbook(sum_file)
    assert wb.sheetnames == ["1班", "2班", "全部班级", "得分率汇总"]
    ws = wb["1班"]
    assert ws.max_row == 25 and ws.max_column == 12
    assert ws.freeze_panes == "B3"
    assert ws.oddHeader.right.text == "2026年05月20日"
    assert "G1:L1" in [str(m) for m in ws.merged_cells.ranges]
    assert ws["D3"].number_format == "0.00%"
    assert ws["E3"].number_format == "0.00"
    cf_types = [r.type for c in ws.conditional_formatting for r in c.rules]
    assert "colorScale" in cf_types
    assert "dataBar" in cf_types

    summary = wb["得分率汇总"]
    assert summary.max_row == 25
    assert summary.max_column == 5  # 题号 + 题型 + 1班 + 2班 + 全部班级
    src = analyze.read_class_sheet(INPUT_DIR / "1班.xls")
    assert abs(summary["C2"].value - src.rows[0][3]) < 1e-9
    expected_sum = sum(
        v for v in (row[4] for row in src.rows) if isinstance(v, (int, float))
    )
    assert abs(summary["C25"].value - expected_sum) < 1e-6

    wb2 = load_workbook(dev_file)
    assert wb2.sheetnames == ["Sheet", "距平统计"]
    ws2 = wb2["Sheet"]
    assert ws2.max_row == 25 and ws2.max_column == 6
    assert ws2["A1"].value is None
    assert ws2["B1"].value == "1班"
    assert abs(ws2["E2"].value - (summary["C2"].value - summary["E2"].value)) < 1e-9
    assert "E2:E24" in [str(c.sqref) for c in ws2.conditional_formatting]

    stat = wb2["距平统计"]
    assert stat["A1"].value is None
    assert stat["B1"].value == "1班_d"
    assert stat["A2"].value == "正距平题数"
    assert stat["A13"].value == "总得分距平"
    assert stat["C13"].value is not None


def test_run_auto_loads_plugin_config(tmp_path, monkeypatch):
    global_cfg, _plugin_cfg, out_dir = _write_configs(
        tmp_path, exam_date="2026-05-20"
    )
    default_cfg = tmp_path / "default_plugin.yaml"
    default_cfg.write_text(
        "input_dir: ''\n"
        "output_dir: ''\n"
        "semester: ''\n"
        "subject: ''\n"
        "exam_date: '2026-05-20'\n"
        "low_score_flag: 0.6\n"
        "max_date_input_errors: 2\n"
        "output_subdir_by_exam: true\n"
        "summary_groups: [teacher]\n"
        "deviation_groups: [teacher]\n"
        "baseline: ''\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        analyze, "_default_plugin_config_path", lambda: str(default_cfg)
    )
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=None,
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 0
    )
    assert _sum_named(out_dir, "柯").is_file()
    assert not _sum_path(out_dir).exists()


def test_summary_group_workbooks(tmp_path):
    global_cfg, plugin_cfg, out_dir = _write_configs(
        tmp_path, exam_date="2026-05-20", summary_groups="[teacher]"
    )
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 0
    )
    wb = load_workbook(_sum_named(out_dir, "柯"))
    assert wb.sheetnames == ["1班", "2班", "全部班级", "得分率汇总"]
    summary = wb["得分率汇总"]
    assert summary.max_column == 5
    assert summary["D2"].border.right.style == "thin"
    assert summary["E2"].border.left.style == "thin"
    assert not _sum_path(out_dir).exists()


def test_baseline_cli_overrides_config(tmp_path):
    global_cfg, plugin_cfg, out_dir = _write_configs(
        tmp_path,
        exam_date="2026-05-20",
        deviation_groups="[level]",
        baseline="柯",
    )
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
            baseline="全部班级",
        )
        == 0
    )
    ws_a = load_workbook(_dev_named(out_dir, "A"))["Sheet"]
    assert ws_a["C1"].value == "全部班级"


def test_deviation_level_files(tmp_path):
    global_cfg, plugin_cfg, out_dir = _write_configs(
        tmp_path, exam_date="2026-05-20", deviation_groups="[level]"
    )
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 0
    )
    ws_a = load_workbook(_dev_named(out_dir, "A"))["Sheet"]
    assert ws_a["B1"].value == "1班"
    assert ws_a["C1"].value == "全部班级"
    ws_b = load_workbook(_dev_named(out_dir, "B"))["Sheet"]
    assert ws_b["B1"].value == "2班"


def test_deviation_baseline_group(tmp_path):
    global_cfg, plugin_cfg, out_dir = _write_configs(
        tmp_path,
        exam_date="2026-05-20",
        deviation_groups="[level]",
        baseline="柯",
    )
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 0
    )
    ws_a = load_workbook(_dev_named(out_dir, "A"))["Sheet"]
    assert ws_a["B1"].value == "1班"
    assert ws_a["C1"].value == "柯"
    assert ws_a["D1"].value == "1班_d"


def test_baseline_bl_alias_in_config(tmp_path):
    global_cfg, _plugin_cfg, out_dir = _write_configs(
        tmp_path, exam_date="2026-05-20", deviation_groups="[level]"
    )
    plugin_cfg = tmp_path / "plugin_bl.yaml"
    plugin_cfg.write_text(
        "input_dir: ''\n"
        "output_dir: ''\n"
        "semester: ''\n"
        "subject: ''\n"
        "exam_date: '2026-05-20'\n"
        "low_score_flag: 0.6\n"
        "max_date_input_errors: 2\n"
        "output_subdir_by_exam: true\n"
        "summary_groups: []\n"
        "deviation_groups: [level]\n"
        "baseline: ''\n"
        "bl: '柯'\n",
        encoding="utf-8",
    )
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 0
    )
    ws_a = load_workbook(_dev_named(out_dir, "A"))["Sheet"]
    assert ws_a["C1"].value == "柯"


def test_explicit_exam_date_wins(tmp_path):
    global_cfg, plugin_cfg, out_dir = _write_configs(
        tmp_path, with_exam=True, exam_date="2026-06-01"
    )
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 0
    )
    wb = load_workbook(_sum_path(out_dir))
    assert wb["1班"].oddHeader.right.text == "2026年06月01日"


def test_find_exam_date_from_config(tmp_path):
    global_cfg, _plugin_cfg, _out = _write_configs(tmp_path, with_exam=True)
    cfg = load_config(str(global_cfg))
    assert (
        analyze.find_exam_date(
            cfg, "高一下地理限时练三", "高一第二学期", "地理", "限时练三"
        )
        == "2026-05-20"
    )


def test_find_exam_date_no_match(tmp_path):
    global_cfg, _plugin_cfg, _out = _write_configs(tmp_path)
    cfg = load_config(str(global_cfg))
    assert (
        analyze.find_exam_date(
            cfg, "高一下地理限时练三", "高一第二学期", "地理", "限时练三"
        )
        is None
    )


def test_find_exam_date_multiple_matches(tmp_path):
    global_cfg, _plugin_cfg, _out = _write_configs(tmp_path)
    cfg = load_config(str(global_cfg))
    (tmp_path / "exams" / "高一第二学期").mkdir()
    (tmp_path / "exams" / "高一第二学期" / "a.yaml").write_text(
        "subject: 地理\n"
        "semester: 高一第二学期\n"
        "date: '2026-05-01'\n"
        "folder: data/input\n"
        "file: a.xlsx\n"
        "name: 限时练三甲\n"
        "short_name: 限时练三\n",
        encoding="utf-8",
    )
    (tmp_path / "exams" / "高一第二学期" / "b.yaml").write_text(
        "subject: 地理\n"
        "semester: 高一第二学期\n"
        "date: '2026-05-02'\n"
        "folder: data/input\n"
        "file: b.xlsx\n"
        "name: 限时练三乙\n"
        "short_name: 限时练三\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="多个匹配"):
        analyze.find_exam_date(
            cfg, "高一下地理限时练三", "高一第二学期", "地理", "限时练三"
        )


def test_interactive_date_input_valid(tmp_path, monkeypatch):
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path)
    answers = iter(["2026/5/20"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 0
    )
    wb = load_workbook(_sum_path(out_dir))
    assert wb["1班"].oddHeader.right.text == "2026年05月20日"


def test_interactive_date_confirm_today(tmp_path, monkeypatch):
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path)
    answers = iter(["bad", "also-bad", "y"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 0
    )
    wb = load_workbook(_sum_path(out_dir))
    assert wb["1班"].oddHeader.right.text == _date.today().strftime("%Y年%m月%d日")


def test_interactive_date_abort(tmp_path, monkeypatch):
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path)
    answers = iter(["bad", "bad2", "n"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 1
    )
    assert not _sum_path(out_dir).exists()


def test_interactive_input_dir_valid(tmp_path, monkeypatch):
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path, exam_date="2026-05-20")
    answers = iter([str(INPUT_DIR)])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=None,
            output_dir=str(out_dir),
        )
        == 0
    )
    assert _sum_path(out_dir).is_file()


def test_interactive_input_dir_abort(tmp_path, monkeypatch):
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path, exam_date="2026-05-20")
    answers = iter(["不存在路径", "bad2"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=None,
            output_dir=str(out_dir),
        )
        == 1
    )
    assert not _sum_path(out_dir).exists()


def test_output_dir_prompt_empty_uses_config(tmp_path, monkeypatch):
    cfg_out = tmp_path / "out_cfg"
    global_cfg, plugin_cfg, _out = _write_configs(
        tmp_path, exam_date="2026-05-20", output_dir_cfg=str(cfg_out)
    )
    answers = iter([""])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=None,
        )
        == 0
    )
    assert _sum_path(cfg_out).is_file()


def test_output_dir_prompt_invalid_uses_config(tmp_path, monkeypatch):
    cfg_out = tmp_path / "out_cfg"
    global_cfg, plugin_cfg, _out = _write_configs(
        tmp_path, exam_date="2026-05-20", output_dir_cfg=str(cfg_out)
    )
    answers = iter([str(tmp_path / "bad<>|:")])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=None,
        )
        == 0
    )
    assert _sum_path(cfg_out).is_file()


def test_output_dir_prompt_valid_custom(tmp_path, monkeypatch):
    user_out = tmp_path / "user_out"
    global_cfg, plugin_cfg, _out = _write_configs(
        tmp_path, exam_date="2026-05-20"
    )
    answers = iter([str(user_out)])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=None,
        )
        == 0
    )
    assert _sum_path(user_out).is_file()


def test_output_dir_no_subdir(tmp_path):
    global_cfg, plugin_cfg, out_dir = _write_configs(
        tmp_path, exam_date="2026-05-20", subdir_by_exam=False
    )
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(INPUT_DIR),
            output_dir=str(out_dir),
        )
        == 0
    )
    assert (out_dir / SUM_NAME).is_file()


def test_run_missing_input_dir(tmp_path):
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path, exam_date="2026-05-20")
    with pytest.raises(FileNotFoundError):
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(tmp_path / "nope"),
            output_dir=str(out_dir),
        )


def test_run_missing_all_class_summary(tmp_path):
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path, exam_date="2026-05-20")
    only_dir = tmp_path / "only"
    only_dir.mkdir()
    (only_dir / "1班.xls").write_bytes((INPUT_DIR / "1班.xls").read_bytes())
    with pytest.raises(ValueError, match="全部班级"):
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(only_dir),
            output_dir=str(out_dir),
        )


def _write_detail_exam_entry(tmp_path: Path, raw_name: str) -> Path:
    """登记一个 file 与小题分原始文件同名的考试条目。"""
    exams_dir = tmp_path / "exams" / "高一第二学期"
    exams_dir.mkdir(parents=True, exist_ok=True)
    entry = exams_dir / "期中联考.yaml"
    entry.write_text(
        "subject: 地理\n"
        "semester: 高一第二学期\n"
        "date: '2026-05-20'\n"
        "folder: .\n"
        f"file: {raw_name}\n"
        "name: 高一下地理期中联考\n"
        "short_name: 期中联考\n",
        encoding="utf-8",
    )
    return entry


def _make_detail_raw(tmp_path: Path) -> Path:
    """构造小题分（含答案）原始文件：答案区 + 得分区。"""
    import pandas as pd

    raw_path = tmp_path / "期中联考地理原始数据.xlsx"
    pd.DataFrame(
        [
            ["姓名", "考号", "学校", "班级", "1", "2", "3", "1", "2", "3"],
            ["张三", "1001", "测试中学", "高一1班", "A", "B", "AB", 2, 2, 3],
            ["李四", "1002", "测试中学", "高一1班", "A", "A", "AC", 2, 0, 3],
            ["王五", "1003", "测试中学", "高一2班", "B", "B", "AB", 2, 2, 3],
            ["赵六", "1004", "测试中学", "高一2班", "A", "A", "AC", 0, 2, 0],
        ]
    ).to_excel(raw_path, index=False, header=False)
    return raw_path


def _detail_sum_path(out_dir: Path) -> Path:
    return (
        out_dir
        / "高一下地理期中联考"
        / "高一下地理期中联考_客观题得分汇总（全部班级）.xlsx"
    )


def test_detail_flow_file_input(tmp_path, monkeypatch):
    """input_dir 直接指向小题分原始文件 → 提示选择学校后生成汇总。"""
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path, exam_date="2026-05-20")
    raw = _make_detail_raw(tmp_path)
    _write_detail_exam_entry(tmp_path, raw.name)
    answers = iter(["测试中学"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(raw),
            output_dir=str(out_dir),
        )
        == 0
    )
    sum_file = _detail_sum_path(out_dir)
    assert sum_file.is_file()
    wb = load_workbook(sum_file)
    assert wb.sheetnames == ["1班", "2班", "全部班级", "得分率汇总"]
    ws1 = wb["1班"]
    assert ws1["A3"].value == "第1题"
    assert ws1["B3"].value == "单选题"
    assert ws1["C3"].value == 2
    assert abs(ws1["D3"].value - 1.0) < 1e-9
    assert abs(ws1["E3"].value - 2.0) < 1e-9
    assert ws1["F3"].value == "A"
    assert abs(ws1["G3"].value - 1.0) < 1e-9
    assert ws1["B5"].value == "多选题"
    ws_all = wb["全部班级"]
    assert ws_all["A3"].value == "第1题"
    assert ws_all["F3"].value in (None, "")
    assert abs(ws_all["D3"].value - 0.75) < 1e-9
    ws2 = wb["2班"]
    assert ws2["F4"].value in (None, "")

def test_detail_flow_folder_without_class_files(tmp_path, monkeypatch):
    """文件夹无 N班.xls → 提示原始文件名称与学校后走新流程。"""
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path, exam_date="2026-05-20")
    raw = _make_detail_raw(tmp_path)
    _write_detail_exam_entry(tmp_path, raw.name)
    answers = iter([raw.name, "测试中学"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(raw.parent),
            output_dir=str(out_dir),
        )
        == 0
    )
    assert _detail_sum_path(out_dir).is_file()

def test_detail_build_sheets_unit(tmp_path):
    """build_detail_class_sheets：班级 sheet 行结构、满分/正确答案规则。"""
    global_cfg, _plugin_cfg, _out = _write_configs(tmp_path)
    cfg = load_config(str(global_cfg))
    raw = _make_detail_raw(tmp_path)
    sheets = analyze.build_detail_class_sheets(raw, cfg, "高一第二学期", "地理")
    names = [s.class_name for s in sheets]
    assert names == ["1班", "2班", "全部班级"]
    row1 = sheets[0].rows[0]
    assert row1[:6] == ["第1题", "单选题", 2, 1.0, 2.0, "A"]
    assert abs(row1[6] - 1.0) < 1e-9
    all_rows = sheets[2].rows
    assert all_rows[0][5] in (None, "")
    assert all_rows[1][5] in (None, "")
    assert all_rows[2][1] == "多选题"


def test_detail_school_invalid_retry(tmp_path, monkeypatch):
    """学校不存在时提示重新输入，正确后继续生成。"""
    global_cfg, plugin_cfg, out_dir = _write_configs(tmp_path, exam_date="2026-05-20")
    raw = _make_detail_raw(tmp_path)
    _write_detail_exam_entry(tmp_path, raw.name)
    answers = iter(["不存在的学校", "测试中学"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    assert (
        analyze.run(
            config_path=str(global_cfg),
            plugin_config=str(plugin_cfg),
            input_dir=str(raw),
            output_dir=str(out_dir),
        )
        == 0
    )
    assert _detail_sum_path(out_dir).is_file()
