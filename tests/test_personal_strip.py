"""个人成绩单生成的单元测试。"""

import pandas as pd

from grade_analyzer.config import AnalysisConfig, ExamConfig, OutputConfig, TeacherMap
from grade_analyzer.personal_strip import build_personal_strips
from grade_analyzer.result_config import ResultsConfig


def _exam() -> ExamConfig:
    return ExamConfig(
        name="高一下期中联考",
        short_name="联考",
        semester="高一第二学期",
        subject="地理",
        date="2026-04-20",
        full_score=100.0,
    )


def _score() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "student_id": ["S1", "S2", "S3", "S4"],
            "name": ["", "", "", ""],
            "class_name": ["高一10班", "高一10班", "高一11班", "高一11班"],
            "class_level": ["A", "A", "B", "B"],
            "teacher": ["柯", "柯", "叶", "叶"],
            "total_score": [80.0, 70.0, 90.0, 60.0],
            "objective_score": [50.0, 40.0, 55.0, 30.0],
            "subjective_score": [30.0, 30.0, 35.0, 30.0],
            "单选分": [40.0, 30.0, 44.0, 25.0],
            "多选分": [10.0, 10.0, 11.0, 5.0],
            "班次": [1, 2, 1, 2],
            "校次": [2, 5, 1, 8],
        }
    )


def _questions() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "student_id": ["S1", "S2", "S3", "S4"] * 2,
            "question_id": ["26-1"] * 4 + ["26-2"] * 4,
            "question_type": ["主观"] * 8,
            "score": [3.0, 2.0, 4.0, 1.0, 2.0, 3.0, 1.0, 2.0],
            "full_score": [None] * 8,
        }
    )


def _config(tmp_path) -> AnalysisConfig:
    cfg = AnalysisConfig(results_dir=str(tmp_path / "results"))
    cfg.teacher_maps = {
        ("高一第二学期", "地理"): TeacherMap(
            subject="地理",
            teacher_count=2,
            teacher_names={"A": "柯", "B": "叶"},
            class_teachers={"高一10班": "A", "高一11班": "B"},
        )
    }
    return cfg


def test_build_all_scope(tmp_path):
    paths = build_personal_strips(
        _exam(), _score(), _questions(), ResultsConfig(), _config(tmp_path)
    )
    assert len(paths) == 1
    assert paths[0].endswith("20260420_高一下期中联考_全部班级_个人成绩单.xlsx")


def test_build_teacher_scope(tmp_path):
    cfg = ResultsConfig()
    cfg.personal.scope.mode = "teacher"
    cfg.personal.scope.teachers = ["柯"]
    paths = build_personal_strips(
        _exam(), _score(), _questions(), cfg, _config(tmp_path)
    )
    assert len(paths) == 1
    assert paths[0].endswith("_柯_个人成绩单.xlsx")


def test_build_teacher_scope_letter_code(tmp_path):
    cfg = ResultsConfig()
    cfg.personal.scope.mode = "teacher"
    cfg.personal.scope.teachers = ["A"]  # 字母代号 -> 柯
    paths = build_personal_strips(
        _exam(), _score(), _questions(), cfg, _config(tmp_path)
    )
    assert len(paths) == 1
    assert paths[0].endswith("_柯_个人成绩单.xlsx")


def test_build_teacher_invalid_raises(tmp_path):
    import pytest

    cfg = ResultsConfig()
    cfg.personal.scope.mode = "teacher"
    cfg.personal.scope.teachers = ["Z"]
    with pytest.raises(ValueError, match="教师配置无效"):
        build_personal_strips(
            _exam(), _score(), _questions(), cfg, _config(tmp_path)
        )


def test_build_custom_scope_label(tmp_path):
    cfg = ResultsConfig()
    cfg.personal.scope.mode = "custom"
    cfg.personal.scope.classes = ["高一10班"]
    paths = build_personal_strips(
        _exam(), _score(), _questions(), cfg, _config(tmp_path)
    )
    assert len(paths) == 1
    assert paths[0].endswith("_柯_个人成绩单.xlsx")  # 与柯老师任教班级一致


def test_build_merged_mode(tmp_path):
    exam = _exam()
    exam.question_display = "merged"
    exam.show_big_questions = True
    cfg = ResultsConfig()
    cfg.personal.merged_prefix = "M"
    cfg.personal.big_score_prefix = "S"
    paths = build_personal_strips(
        exam, _score(), _questions(), cfg, _config(tmp_path)
    )
    import openpyxl

    ws = openpyxl.load_workbook(paths[0])["个人成绩单"]
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert "M26" in header  # 合并列：前缀+大题题号
    assert "S26" in header  # 大题总分列：前缀+大题题号
    assert "26(1)" not in header
    merged = ws.cell(2, header.index("M26") + 1).value
    assert merged == "3|2"  # 按[班级,总分,考号]排序后首位为 高一10班 S1
    assert ws.cell(2, header.index("S26") + 1).value == 5.0  # 3+2
    assert ws.column_dimensions["K"].width == 8  # 合并小题列宽
    assert ws.column_dimensions["L"].width == 4  # 大题列宽


def test_strip_style(tmp_path):
    import openpyxl

    paths = build_personal_strips(
        _exam(), _score(), _questions(), ResultsConfig(), _config(tmp_path)
    )
    ws = openpyxl.load_workbook(paths[0])["个人成绩单"]
    assert ws.cell(1, 1).font.name == "宋体"
    assert ws.cell(1, 1).border.top.style == "medium"  # 表头上框线略粗
    assert ws.cell(2, 1).border.top.style == "thin"  # 数据行正常
    assert ws.cell(2, 6).alignment.horizontal == "center"  # 总分居中
    assert ws.cell(2, 7).alignment.horizontal == "right"  # 客观分右对齐
    assert ws.cell(2, 1).alignment.vertical == "center"  # 垂直居中


def test_strip_sorted_by_class_score_id(tmp_path):
    import openpyxl

    paths = build_personal_strips(
        _exam(), _score(), _questions(), ResultsConfig(), _config(tmp_path)
    )
    ws = openpyxl.load_workbook(paths[0])["个人成绩单"]
    assert ws.cell(2, 1).value == "高一10班"  # 第一块为高一10班
    assert ws.cell(2, 6).value == 80.0  # S1 总分（降序）
    assert ws.cell(5, 1).value == "高一10班"  # 第二块仍为高一10班（S2）
    assert ws.cell(5, 6).value == 70.0
    assert ws.cell(8, 1).value == "高一11班"  # 第三块为高一11班
    assert ws.cell(8, 6).value == 90.0  # S3 总分


def test_strip_widths_and_bold(tmp_path):
    import openpyxl

    paths = build_personal_strips(
        _exam(), _score(), _questions(), ResultsConfig(), _config(tmp_path)
    )
    ws = openpyxl.load_workbook(paths[0])["个人成绩单"]
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert ws.column_dimensions["A"].width == 10  # 班级
    assert ws.column_dimensions["J"].width == 6  # 多选（前10列末列）
    sub_idx = header.index("26(1)") + 1
    assert ws.column_dimensions[chr(64 + sub_idx)].width == 4  # 单独小题列宽
    assert ws.cell(2, 6).font.bold is True  # 总分加粗
    assert ws.cell(2, 1).font.bold is False


def _exam2() -> ExamConfig:
    return ExamConfig(
        name="高一下地理限时练一",
        short_name="限时练一",
        semester="高一第二学期",
        subject="地理",
        date="2026-03-25",
        full_score=100.0,
    )


def _score2() -> pd.DataFrame:
    # 第二场：S1、S2 参加（S3、S4 未参加）
    df = _score().copy()
    df = df[df["student_id"].isin(["S1", "S2"])]
    df["total_score"] = [85.0, 65.0]
    df["班次"] = [1, 2]
    df["校次"] = [1, 6]
    return df


def test_build_merged_personal_strips(tmp_path):
    import openpyxl

    from grade_analyzer.personal_strip import build_merged_personal_strips

    paths = build_merged_personal_strips(
        [_exam(), _exam2()],
        [_score(), _score2()],
        [_questions(), _questions()],
        ResultsConfig(),
        _config(tmp_path),
    )
    assert len(paths) == 1
    assert paths[0].endswith("20260325-20260420_全部班级_个人成绩单.xlsx")

    ws = openpyxl.load_workbook(paths[0])["个人成绩单"]
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert header == [
        "班级", "姓名", "考试", "班次", "校次",
        "总分", "客观分", "主观分", "单选", "多选", "主观题",
    ]
    # 排序：班级10 中 S1 平均分 82.5 > S2 67.5 → S1 块在前
    assert ws.cell(1, 1).value == "班级"  # S1 的表头
    assert ws.cell(2, 3).value == "联考"  # S1 第 1 场（2026-04-20）
    assert ws.cell(2, 6).value == 80.0
    assert ws.cell(3, 3).value == "限时练一"  # S1 第 2 场
    assert ws.cell(3, 6).value == 85.0
    assert ws.cell(2, 11).value == "5"  # 主观大题 26 得分（3+2）合并
    # S1 块：表头 + 2 行 + 空行(blank_rows_between=1)；S2 块从第 5 行开始
    assert ws.cell(5, 1).value == "班级"
    assert ws.cell(6, 3).value == "联考"
    # 默认容量（37 行）下 4 个学生块（14 行）不产生分页符
    assert len(ws.row_breaks.brk) == 0


def test_rows_per_page_default():
    from grade_analyzer.personal_strip import _rows_per_page

    assert _rows_per_page(ResultsConfig()) == 37  # (11.69-1)*72/20 向下取整后减 1 保守余量


def test_rows_per_page_config_priority():
    from grade_analyzer.personal_strip import _rows_per_page

    cfg = ResultsConfig()
    cfg.personal.print.rows_per_page = 41
    assert _rows_per_page(cfg) == 41  # 配置优先


def test_apply_page_breaks_inserts_before_headers():
    from openpyxl import Workbook

    from grade_analyzer.personal_strip import _apply_page_breaks

    cfg = ResultsConfig()
    cfg.personal.layout.row_height = 1000  # 页容量 1 行
    ws = Workbook().active
    _apply_page_breaks(ws, {1, 4, 7}, {1: 3, 4: 3, 7: 3}, cfg)
    assert [b.id for b in ws.row_breaks.brk] == [3, 6]  # 0-based：表头 4/7 前分页


def test_apply_page_breaks_no_break_when_fits():
    from openpyxl import Workbook

    from grade_analyzer.personal_strip import _apply_page_breaks

    ws = Workbook().active
    _apply_page_breaks(ws, {1, 4, 7}, {1: 3, 4: 3, 7: 3}, ResultsConfig())
    assert len(ws.row_breaks.brk) == 0


def test_strip_page_breaks_inserted(tmp_path):
    import openpyxl

    cfg = ResultsConfig()
    cfg.personal.layout.row_height = 1000
    paths = build_personal_strips(
        _exam(), _score(), _questions(), cfg, _config(tmp_path)
    )
    wb = openpyxl.load_workbook(paths[0])
    ws = wb[wb.sheetnames[0]]
    ids = [b.id for b in ws.row_breaks.brk]
    assert {3, 6, 9} <= set(ids)  # 每个学生块表头前分页（0-based）


def test_merged_page_breaks_inserted(tmp_path):
    import openpyxl

    from grade_analyzer.personal_strip import build_merged_personal_strips

    cfg = ResultsConfig()
    cfg.personal.layout.row_height = 1000
    paths = build_merged_personal_strips(
        [_exam(), _exam2()],
        [_score(), _score2()],
        [_questions(), _questions()],
        cfg,
        _config(tmp_path),
    )
    wb = openpyxl.load_workbook(paths[0])
    ws = wb[wb.sheetnames[0]]
    assert {4, 8, 11} <= {b.id for b in ws.row_breaks.brk}  # 每块表头前分页（0-based）


def test_build_strip_rows_drops_blank_when_full_page():
    from grade_analyzer.personal_strip import _build_strip_rows
    from grade_analyzer.report import _subjective_pivot

    cfg = ResultsConfig()
    cfg.personal.print.rows_per_page = 8
    wide, big, subj_cols, big_cols = _subjective_pivot(_questions())
    cols, rows, widths, header_rows, block_sizes = _build_strip_rows(
        _score(), _exam(), wide, big, subj_cols, big_cols,
        {"高一10班", "高一11班"}, cfg,
    )
    # 容量 8：S1/S2 各 3 行后 used=6，S3（表头 7）删空行后 2 行刚好满页
    assert block_sizes[7] == 2
    assert 9 in header_rows  # S4 表头紧随（无空行间隔）
    assert rows[8] == list(cols)


def test_build_merged_rows_drops_blank_when_full_page():
    from grade_analyzer.personal_strip import _build_merged_rows

    cfg = ResultsConfig()
    cfg.personal.print.rows_per_page = 7
    valid_list = [_score().copy(), _score2().copy()]
    cols, rows, widths, header_rows, block_sizes = _build_merged_rows(
        [_exam(), _exam2()],
        valid_list,
        [_questions(), _questions()],
        {"高一10班", "高一11班"},
        cfg,
    )
    # 容量 7：S1 块 4 行后 used=4，S2（表头 5）删空行 3 行刚好满页
    assert block_sizes[5] == 3
    assert 8 in header_rows  # S3 表头紧随


def test_merged_rows_sorted_by_exam_count_group():
    from grade_analyzer.personal_strip import _build_merged_rows

    def _mk(students):
        rows = []
        for sid, cls, total in students:
            rows.append(
                {
                    "student_id": sid, "name": "", "class_name": cls,
                    "class_level": "A", "teacher": "",
                    "total_score": total, "objective_score": 0.0,
                    "subjective_score": 0.0, "单选分": 0.0, "多选分": 0.0,
                    "班次": 1, "校次": 1,
                }
            )
        return pd.DataFrame(rows)

    # A 两场（班级11）、B 一场（班级10）、C 两场（班级10）
    s1 = _mk([("A", "高一11班", 80.0), ("B", "高一10班", 90.0), ("C", "高一10班", 70.0)])
    s2 = _mk([("A", "高一11班", 85.0), ("C", "高一10班", 75.0)])
    cols, rows, widths, header_rows, block_sizes = _build_merged_rows(
        [_exam(), _exam2()],
        [s1, s2],
        [_questions(), _questions()],
        {"高一10班", "高一11班"},
        ResultsConfig(),
    )
    # 场次数量优先：C(2场,班10)、A(2场,班11)、B(1场,班10)
    assert rows[1][0] == "高一10班"  # C 第一行数据
    assert rows[5][0] == "高一11班"  # A 第一行数据
    assert rows[9][0] == "高一10班"  # B 数据（一场）
