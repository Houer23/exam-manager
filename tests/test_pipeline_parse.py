"""parse 流程的集成测试。"""

from pathlib import Path

import pandas as pd
import pytest
import yaml

from grade_analyzer.config import load_config
from grade_analyzer.pipeline import parse_exams


def _sample_joint_school() -> str:
    """读取本地样例联考文件的学校名（样例数据不入库，测试中不写死真实校名）。"""
    from grade_analyzer.io_utils import read_raw_sheet

    raw = read_raw_sheet("data/input/测试样例/地理原始数据.xlsx")
    for r in range(min(6, len(raw))):
        for j, v in enumerate(raw.iloc[r].tolist()):
            if str(v).strip() == "学校":
                value = raw.iloc[r + 1, j]
                return str(value).strip() if value is not None else ""
    return "示例二中"


def _setup(tmp_path, parsed_dir, out_dir=None, roster_dir=None, default_school="示例二中"):
    exams_dir = tmp_path / "exams"
    classes_dir = tmp_path / "classes"
    subjects_dir = tmp_path / "subjects"
    (exams_dir / "高一第二学期").mkdir(parents=True)
    classes_dir.mkdir(exist_ok=True)
    subjects_dir.mkdir(exist_ok=True)
    (exams_dir / "高一第二学期" / "周测.yaml").write_text(
        "name: 周测\n"
        "format: weekly\n"
        "date: \"2026-03-25\"\n"
        "folder: data/input/测试样例\n"
        "file: 【教学班报告--高一下地理限时练一】所有班级学生小题得分明细.xlsx\n"
        "subject: 地理\n"
        "full_score: 100\n"
        "objective_full_score: 85\n"
        "subjective_full_score: 15\n",
        encoding="utf-8",
    )
    # 测试用班级/学科配置：不依赖工作区真实配置，新克隆环境也可运行
    (classes_dir / "高一第二学期.yaml").write_text(
        "高一13班: {level: A, course: 物化地}\n"
        "高一6班: {level: B, course: 史地政}\n"
        "高一16班: {level: A, course: 物化地}\n",
        encoding="utf-8",
    )
    (subjects_dir / "高一第二学期_地理.yaml").write_text(
        "subject: 地理\n"
        "teacher_count: 1\n"
        "teacher_names:\n"
        "  A: 柯\n"
        "class_teachers:\n"
        "  高一13班: A\n",
        encoding="utf-8",
    )
    cfg = tmp_path / "config.yaml"
    data = {
        "subjects": ["语文", "数学", "外语", "物理", "化学", "生物", "政治", "历史", "地理", "技术"],
        "exams_dir": str(exams_dir),
        "classes_dir": str(classes_dir),
        "subjects_dir": str(subjects_dir),
        "parsed_dir": str(parsed_dir),
        "default_school": default_school,
        "output_dir": str(out_dir or tmp_path / "out"),
        "report_excel_name": "成绩分析汇总.xlsx",
        "results_dir": str((out_dir or tmp_path / "out") / "results"),
    }
    if roster_dir:
        data["roster_dir"] = str(roster_dir)
    results_cfg_dir = tmp_path / "results_cfg"
    results_cfg_dir.mkdir(exist_ok=True)
    (results_cfg_dir / "config.yaml").write_text(
        "personal:\n  scope:\n    mode: all\n", encoding="utf-8"
    )
    data["results_config_dir"] = str(results_cfg_dir)
    cfg.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    return str(cfg)


@pytest.fixture(scope="module")
def shared_parsed(tmp_path_factory):
    """模块内共享：解析一次真实周测样例，供多个测试复用规范表，减少重复解析。"""
    from grade_analyzer.pipeline import parse_exams

    base = tmp_path_factory.mktemp("shared_parsed")
    parsed_dir = base / "parsed"
    cfg_path = _setup(base, parsed_dir, base / "out")
    parse_exams(cfg_path)
    return parsed_dir


def test_run_pipeline_verify_roster(shared_parsed, tmp_path, capsys):
    from grade_analyzer.pipeline import run_pipeline

    parsed = shared_parsed
    out = tmp_path / "out"
    roster_dir = tmp_path / "roster"
    roster_dir.mkdir()
    pd.DataFrame(
        {
            "姓名": ["柯一"],
            "考号": ["250907010858"],
            "性别": ["男"],
            "七选三": ["物化地"],
            "班级": ["高一年级13班"],
        }
    ).to_excel(roster_dir / "高一第二学期.xlsx", index=False)

    cfg_path = _setup(tmp_path, parsed, out, str(roster_dir))
    run_pipeline(cfg_path, verify_roster=True)
    out_text = capsys.readouterr().out
    assert "[名单核对]" in out_text
    assert (out / "quality" / "名单核对_高一下地理周测.csv").is_file()


def test_parse_exams_writes_and_reuses(tmp_path, capsys):
    parsed = tmp_path / "parsed"
    cfg_path = _setup(tmp_path, parsed, tmp_path / "out")

    parse_exams(cfg_path)
    out1 = capsys.readouterr().out
    assert "[完成]" in out1
    assert "高一下地理周测" in out1

    score_path = parsed / "高一第二学期" / "高一下地理周测" / "score_summary.csv"
    question_path = parsed / "高一第二学期" / "高一下地理周测" / "question_detail.csv"
    assert score_path.is_file()
    assert question_path.is_file()
    score_df = pd.read_csv(score_path, encoding="utf-8-sig")
    assert len(score_df) > 400
    assert (score_df["school"] == "示例二中").all()
    assert (score_df["grade"] == "高一").all()
    assert (score_df["class_name"] == "高一13班").any()
    assert {"班次", "校次"} <= set(score_df.columns)
    assert score_df["班次"].dropna().min() == 1.0
    assert {"单选分", "多选分", "单选满分", "多选满分"} <= set(score_df.columns)
    assert {"class_level", "course", "teacher"} <= set(score_df.columns)
    assert "exam_date" in score_df.columns
    assert set(score_df["class_level"].dropna().unique()) <= {"A", "B"}
    assert score_df["单选满分"].iloc[0] == 60.0  # 周测：3 × 20
    assert score_df["多选满分"].iloc[0] == 25.0  # 周测：5 × 5
    assert len(score_df) * 29 == len(
        pd.read_csv(question_path, encoding="utf-8-sig")
    )
    q_df = pd.read_csv(question_path, encoding="utf-8-sig")
    type_counts = q_df["question_type"].value_counts().to_dict()
    assert type_counts.get("单选") == len(score_df) * 20
    assert type_counts.get("多选") == len(score_df) * 5
    assert type_counts.get("主观") == len(score_df) * 4

    parse_exams(cfg_path)
    assert "[复用]" in capsys.readouterr().out

    parse_exams(cfg_path, reparse=True)
    assert "[复用]" not in capsys.readouterr().out


def test_parse_with_objective_question_count(tmp_path):
    """客观题数配置后：题号大于该数的纯数字题也认定为主观题。"""
    from grade_analyzer.config import load_config
    from grade_analyzer.pipeline import parse_exams
    from grade_analyzer.storage import read_question_detail

    parsed = tmp_path / "parsed"
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)
    exam_yaml = tmp_path / "exams" / "高一第二学期" / "周测.yaml"
    text = exam_yaml.read_text(encoding="utf-8")
    exam_yaml.write_text(text + "objective_question_count: 24\n", encoding="utf-8")

    parse_exams(cfg_path)
    cfg = load_config(cfg_path)
    q = read_question_detail(str(parsed), cfg.exams[0], "csv")
    subj_ids = set(q[q["question_type"] == "主观"]["question_id"])
    assert "25" in subj_ids  # 题号 25 > 24，原客观题变主观
    assert "26-1" in subj_ids


def test_run_pipeline_writes_report_and_statistics(shared_parsed, tmp_path, capsys):
    import openpyxl
    from grade_analyzer.pipeline import run_pipeline

    parsed = shared_parsed
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)

    run_pipeline(cfg_path)
    out_text = capsys.readouterr().out
    assert "[统计]" in out_text
    assert "[报告]" in out_text

    report = out / "reports" / "高一第二学期" / "成绩分析汇总_20260325.xlsx"
    assert report.is_file()
    stat = out / "statistics" / "高一第二学期" / "高一下地理周测.xlsx"
    assert stat.is_file()
    sheets = pd.read_excel(stat, sheet_name=None)
    assert set(sheets) == {"科目统计", "分数段分布", "个人排名", "班级对比", "教师对比"}
    assert not (out / "merged" / "merged_long.csv").exists()  # 单场不做 merge 落盘
    assert (out / "run-info" / "latest" / "运行日志.txt").is_file()
    assert (out / "quality" / "数据质量.xlsx").is_file()
    results = out / "results" / "高一第二学期"
    assert (results / "高一下地理周测" / "高一下地理周测_柯_班级成绩汇总.xlsx").is_file()
    assert (results / "高一下地理周测" / "高一下地理周测_全部班级_班级成绩汇总.xlsx").is_file()
    assert list(results.glob("*_全部班级_个人成绩单.xlsx"))

    wb = openpyxl.load_workbook(stat)
    ws = wb["科目统计"]
    header = {cell.value: cell.column for cell in ws[1] if cell.value}
    assert ws.cell(row=2, column=header["平均分"]).number_format == "0.0"
    assert ws.cell(row=2, column=header["及格率"]).number_format == "0.00%"
    assert len(wb["分数段分布"]._charts) == 1  # 内嵌柱状图
    assert len(wb["班级对比"]._charts) == 1  # 内嵌条形图
    report_wb = openpyxl.load_workbook(report)
    assert len(report_wb["多场趋势"]._charts) == 1  # 内嵌折线图


def test_run_results_exam_filter_with_unnamed_weekly(shared_parsed, tmp_path, capsys):
    """name 留空的周测：--exam 应能用文件名解析出的名称命中。"""
    from grade_analyzer.pipeline import parse_exams, run_results

    parsed = shared_parsed
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)
    # 模拟真实周测：name 留空，考试名称从文件名提取
    (tmp_path / "exams" / "高一第二学期" / "周测.yaml").write_text(
        "name: ''\n"
        "format: weekly\n"
        "folder: data/input/测试样例\n"
        "file: 【教学班报告--高一下地理限时练一】所有班级学生小题得分明细.xlsx\n"
        "subject: 地理\n"
        "full_score: 100\n"
        "objective_full_score: 85\n"
        "subjective_full_score: 15\n",
        encoding="utf-8",
    )

    parse_exams(cfg_path)
    run_results(cfg_path, exam_name="高一下地理限时练一")
    out_text = capsys.readouterr().out
    assert "[班级汇总]" in out_text
    assert "[个人成绩单]" in out_text


def test_run_results_current_exam_index_lists_exams(shared_parsed, tmp_path, capsys):
    """未传 --exam 时列出考试列表，并按 current_exam 序号限定。"""
    from grade_analyzer.pipeline import parse_exams, run_results

    parsed = shared_parsed
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)
    parse_exams(cfg_path)
    # 设置 current_exam = 1（第 1 场）
    text = Path(cfg_path).read_text(encoding="utf-8")
    Path(cfg_path).write_text(text + "current_exam: 1\n", encoding="utf-8")

    run_results(cfg_path)
    out_text = capsys.readouterr().out
    assert "[考试列表]" in out_text
    assert "1. 高一下地理周测" in out_text
    assert "[个人成绩单]" in out_text


def test_run_pipeline_exam_filter(shared_parsed, tmp_path, capsys):
    """run --exam 只处理指定考试；未指定时含无规范表考试会失败。"""
    from grade_analyzer.pipeline import run_pipeline

    parsed = shared_parsed
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)
    # 追加一场无规范表、原始文件也不存在的考试
    (tmp_path / "exams" / "高一第二学期" / "缺考联考.yaml").write_text(
        "name: 高一下缺考联考\n"
        "format: joint\n"
        "folder: data/input\n"
        "file: 不存在.xlsx\n"
        "subject: 地理\n",
        encoding="utf-8",
    )

    run_pipeline(cfg_path, exam="高一下地理周测")
    out_text = capsys.readouterr().out
    assert "[统计] 高一下地理周测" in out_text
    assert "高一下地理缺考联考" not in out_text

    with pytest.raises(ValueError, match="未找到指定考试"):
        run_pipeline(cfg_path, exam="不存在的考试")


def test_resolve_exam_selection_number_list():
    from grade_analyzer.config import AnalysisConfig, ExamConfig
    from grade_analyzer.pipeline import _resolve_exam_selection

    cfg = AnalysisConfig(current_exam=None)
    cfg.exams = [
        ExamConfig(name="A", semester="高一第二学期", date="2026-03-25"),
        ExamConfig(name="B", semester="高一第二学期", date="2026-04-20"),
        ExamConfig(name="C", semester="高一第二学期", date="2026-05-01"),
    ]
    selected, names, merge_semester = _resolve_exam_selection(
        cfg, "1,3", "高一第二学期"
    )
    assert [e.name for e in selected] == ["A", "C"]
    assert names == ["A", "C"]
    assert merge_semester == "高一第二学期"

    selected2, _, _ = _resolve_exam_selection(cfg, "B", "高一第二学期")
    assert [e.name for e in selected2] == ["B"]


def test_resolve_exam_selection_current_exam_list():
    from grade_analyzer.config import AnalysisConfig, ExamConfig
    from grade_analyzer.pipeline import _resolve_exam_selection

    cfg = AnalysisConfig(current_exam=[1, -1])
    cfg.exams = [
        ExamConfig(name="A", semester="高一第二学期", date="2026-03-25"),
        ExamConfig(name="B", semester="高一第二学期", date="2026-04-20"),
        ExamConfig(name="C", semester="高一第二学期", date="2026-05-01"),
    ]
    selected, names, _ = _resolve_exam_selection(cfg, None, "高一第二学期")
    assert [e.name for e in selected] == ["A", "C"]  # 第 1 场与倒数第 1 场
    assert names == ["A", "C"]


def test_cli_results_no_merge_strips_parses():
    from grade_analyzer.cli import build_parser

    args = build_parser().parse_args(["results", "--no-merge-strips"])
    assert args.no_merge_strips is True
    args2 = build_parser().parse_args(["results"])
    assert args2.no_merge_strips is False


def test_cli_results_only_flags_mutually_exclusive():
    from grade_analyzer.cli import build_parser

    args = build_parser().parse_args(["results", "--summary-only"])
    assert args.summary_only is True
    assert args.strips_only is False
    args2 = build_parser().parse_args(["results", "--strips-only"])
    assert args2.summary_only is False
    assert args2.strips_only is True
    with pytest.raises(SystemExit):
        build_parser().parse_args(["results", "--summary-only", "--strips-only"])


def test_run_results_merge_strips_false_generates_per_exam(shared_parsed, tmp_path, capsys):
    """merge_strips=False 时多场各自生成单场个人成绩单，不生成范围合并文件。"""
    from grade_analyzer.pipeline import parse_exams, run_results

    parsed = shared_parsed
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out, default_school=_sample_joint_school())
    # 追加第二场联考条目（真实样例）
    (tmp_path / "exams" / "高一第二学期" / "联考.yaml").write_text(
        "name: 高一下地理期中联考\n"
        "format: joint\n"
        "date: \"2026-04-20\"\n"
        "folder: data/input/测试样例\n"
        "file: 地理原始数据.xlsx\n"
        "subject: 地理\n"
        "full_score: 100\n"
        "objective_full_score: 55\n"
        "subjective_full_score: 45\n",
        encoding="utf-8",
    )
    parse_exams(cfg_path)

    run_results(cfg_path, merge_strips=False)
    out_text = capsys.readouterr().out
    assert "[个人成绩单]" in out_text
    results = out / "results" / "高一第二学期"
    assert not list(results.glob("*-*_*_个人成绩单.xlsx"))  # 无范围合并文件
    assert list(results.glob("20260325_*_个人成绩单.xlsx"))
    assert list(results.glob("20260420_*_个人成绩单.xlsx"))


def test_run_results_summary_only(shared_parsed, tmp_path, capsys):
    """generate_strips=False：只生成班级汇总，不生成个人成绩单。"""
    from grade_analyzer.pipeline import parse_exams, run_results

    parsed = shared_parsed
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)
    parse_exams(cfg_path)
    run_results(cfg_path, generate_strips=False)
    out_text = capsys.readouterr().out
    assert "[班级汇总]" in out_text
    assert "[个人成绩单]" not in out_text


def test_run_results_strips_only(shared_parsed, tmp_path, capsys):
    """generate_summary=False：只生成个人成绩单，不生成班级汇总。"""
    from grade_analyzer.pipeline import parse_exams, run_results

    parsed = shared_parsed
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)
    parse_exams(cfg_path)
    run_results(cfg_path, generate_summary=False)
    out_text = capsys.readouterr().out
    assert "[班级汇总]" not in out_text
    assert "[个人成绩单]" in out_text


def test_ensure_parsed_ready(tmp_path):
    from grade_analyzer.config import load_config
    from grade_analyzer.pipeline import _ensure_parsed_ready, parse_exams

    parsed = tmp_path / "parsed"
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)
    cfg = load_config(cfg_path)
    issues = _ensure_parsed_ready(cfg, cfg.exams)
    assert any("未解析" in i for i in issues)

    parse_exams(cfg_path)
    cfg2 = load_config(cfg_path)
    assert _ensure_parsed_ready(cfg2, cfg2.exams) == []


def test_run_results_requires_parse(tmp_path, capsys):
    """未 parse 时 results 提示先 parse 并正常结束（不报错）。"""
    from grade_analyzer.pipeline import run_results

    parsed = tmp_path / "parsed"
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)
    run_results(cfg_path)
    out_text = capsys.readouterr().out
    assert "请先运行 parse" in out_text
    assert "[个人成绩单]" not in out_text  # 未继续生成


def test_cli_charts_parses():
    from grade_analyzer.cli import build_parser

    args = build_parser().parse_args(["charts", "--exam", "1,3"])
    assert args.command == "charts"
    assert args.exam == "1,3"
    args2 = build_parser().parse_args(["charts"])
    assert args2.exam is None


def test_run_charts_generates(shared_parsed, tmp_path, capsys):
    """已 parse 后 charts 命令生成统计图。"""
    from grade_analyzer.pipeline import run_charts

    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, shared_parsed, out)
    run_charts(cfg_path)
    out_text = capsys.readouterr().out
    assert "[统计图]" in out_text
    charts = out / "charts" / "高一第二学期"
    assert list(charts.glob("按层次_*.png"))
    assert list(charts.glob("按教师_*.png"))


def test_run_charts_requires_parse(tmp_path, capsys):
    """未 parse 时 charts 提示先 parse 并正常结束。"""
    from grade_analyzer.pipeline import run_charts

    parsed = tmp_path / "parsed"
    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, parsed, out)
    run_charts(cfg_path)
    out_text = capsys.readouterr().out
    assert "请先运行 parse" in out_text
    assert "[统计图]" not in out_text


def test_resolve_class_names():
    from grade_analyzer.pipeline import _resolve_class_names

    assert _resolve_class_names("10,11", "高一") == ["高一10班", "高一11班"]
    assert _resolve_class_names("3", "高二") == ["高二03班"]
    assert _resolve_class_names("10-12", "高一") == [
        "高一10班", "高一11班", "高一12班",
    ]
    assert _resolve_class_names("12-10", "高一") == [
        "高一12班", "高一11班", "高一10班",
    ]
    assert _resolve_class_names("10,13-14,10", "高一") == [
        "高一10班", "高一13班", "高一14班",
    ]
    with pytest.raises(ValueError, match="班级"):
        _resolve_class_names("abc", "高一")
    with pytest.raises(ValueError, match="班级"):
        _resolve_class_names("a-b", "高一")


def test_run_charts_with_class(shared_parsed, tmp_path, capsys):
    """charts --class 按 班级×考试 绘制。"""
    from grade_analyzer.pipeline import run_charts

    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, shared_parsed, out)
    run_charts(cfg_path, classes="13")  # 高一13班
    out_text = capsys.readouterr().out
    assert "[统计图]" in out_text
    charts = out / "charts" / "高一第二学期"
    assert list(charts.glob("按班级_*.png"))


def test_run_charts_per_class(shared_parsed, tmp_path, capsys):
    """--per-class 时每个班级单独生成一张图。"""
    from grade_analyzer.pipeline import run_charts

    out = tmp_path / "out"
    cfg_path = _setup(tmp_path, shared_parsed, out)
    run_charts(cfg_path, classes="13,14", per_class=True)
    out_text = capsys.readouterr().out
    assert out_text.count("[统计图]") == 2
    charts = out / "charts" / "高一第二学期"
    assert list(charts.glob("按班级_20260325_高一13班.png"))
    assert list(charts.glob("按班级_20260325_高一14班.png"))
