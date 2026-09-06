"""班级成绩汇总的单元测试。"""

from pathlib import Path

import openpyxl
import pandas as pd

from grade_analyzer.config import AnalysisConfig, ExamConfig, TeacherMap
from grade_analyzer.report import _subjective_pivot, build_class_summaries
from grade_analyzer.result_config import ResultsConfig


def test_subjective_pivot():
    questions = pd.DataFrame(
        {
            "student_id": ["S1", "S1", "S2", "S2", "S1", "S2"],
            "question_id": ["26-1", "26-2", "26-1", "26-2", "27-1", "27-1"],
            "question_type": ["主观"] * 6,
            "score": [3.0, 2.0, 4.0, 1.0, 5.0, 6.0],
        }
    )
    wide, big, subj_cols, big_cols = _subjective_pivot(questions)
    assert subj_cols == ["26-1", "26-2", "27-1"]
    assert big_cols == ["26", "27"]


def test_subjective_pivot_numeric_ids():
    """主观题号为纯数字（int）时（配置客观题数后）也能正常透视。"""
    questions = pd.DataFrame(
        {
            "student_id": ["S1", "S1", "S2", "S2"],
            "question_id": [21, 22, 21, 22],
            "question_type": ["主观"] * 4,
            "score": [3.0, 4.0, 2.0, 5.0],
            "full_score": [None] * 4,
        }
    )
    wide, big, subj_cols, big_cols = _subjective_pivot(questions)
    assert subj_cols == ["21", "22"]
    assert big_cols == ["21", "22"]
    assert wide.loc["S1", "21"] == 3.0
    assert big.loc["S2", "22"] == 5.0


def test_subjective_pivot_keeps_chinese_column_config_order():
    """中文整列大题按 question_detail 行序（即 question_types 书写顺序）展示。"""
    questions = pd.DataFrame(
        {
            "student_id": ["S1"] * 3,
            "question_id": [
                "语法填空56-65", "应用文", "续写",
            ],
            "question_type": ["主观"] * 3,
            "score": [7.5, 6.0, 4.5],
        }
    )
    wide, big, subj_cols, big_cols = _subjective_pivot(questions)
    assert subj_cols == ["语法填空56-65", "应用文", "续写"]
    assert big_cols == ["语法填空56-65", "应用文", "续写"]


def test_subjective_pivot_numeric_before_chinese_columns():
    """数字主观题号在前，中文整列大题按配置顺序排在其后。"""
    questions = pd.DataFrame(
        {
            "student_id": ["S1"] * 5,
            "question_id": [
                "15", "16", "语法填空56-65", "应用文", "续写",
            ],
            "question_type": ["主观"] * 5,
            "score": [3.0, 4.0, 7.5, 6.0, 4.5],
        }
    )
    wide, big, subj_cols, big_cols = _subjective_pivot(questions)
    assert subj_cols == ["15", "16", "语法填空56-65", "应用文", "续写"]
    assert big_cols == ["15", "16", "语法填空56-65", "应用文", "续写"]


def test_build_class_summaries(tmp_path):
    config = AnalysisConfig(
        parsed_dir=str(tmp_path / "parsed"),
        results_dir=str(tmp_path / "results"),
    )
    config.teacher_maps = {
        ("高一第二学期", "地理"): TeacherMap(
            subject="地理",
            teacher_count=1,
            teacher_names={"A": "柯老师"},
            class_teachers={"高一10班": "A", "高一11班": "A"},
        )
    }
    exam = ExamConfig(
        name="测试", semester="高一第二学期", subject="地理", date="2026-04-20"
    )
    score = pd.DataFrame(
        {
            "student_id": ["S1", "S2", "S3"],
            "name": ["甲", "乙", "丙"],
            "class_name": ["高一10班", "高一10班", "高一11班"],
            "total_score": [80.0, 90.0, 70.0],
            "objective_score": [50.0, 55.0, 40.0],
            "subjective_score": [30.0, 35.0, 30.0],
            "单选分": [40.0, 44.0, 30.0],
            "多选分": [10.0, 11.0, 10.0],
            "班次": [2, 1, 1],
            "校次": [2, 1, 3],
            "teacher": ["柯老师", "柯老师", "柯老师"],
        }
    )
    questions = pd.DataFrame(
        {
            "student_id": ["S1", "S1", "S2", "S2", "S3", "S3"],
            "question_id": ["26-1", "26-2", "26-1", "26-2", "26-1", "26-2"],
            "question_type": ["主观"] * 6,
            "score": [3.0, 2.0, 4.0, 1.0, 3.0, 0.0],
        }
    )

    paths = build_class_summaries(exam, score, questions, config, ResultsConfig())
    assert len(paths) == 2  # 教师文件 + 全部班级汇总
    path = Path(paths[0])
    assert path.name == "测试_柯老师_班级成绩汇总.xlsx"
    all_path = Path(paths[1])
    assert all_path.name == "测试_全部班级_班级成绩汇总.xlsx"

    wb = openpyxl.load_workbook(path)
    assert set(wb.sheetnames) == {"高一10班", "高一11班"}
    ws = wb["高一10班"]
    header = [c.value for c in ws[1]]
    assert header[:7] == ["序号", "姓名", "总分", "客观分", "主观分", "单选", "多选"]
    assert "26(1)" in header and "26" in header  # 小题表头为参考样式
    assert "26-1" not in header
    assert header[-2:] == ["校次", "班次"]

    rows = ws.max_row
    assert ws.cell(row=rows - 2, column=1).value == "平均(全班)"
    assert ws.cell(row=rows - 1, column=1).value == "平均(前半)"
    assert ws.cell(row=rows, column=1).value == "平均(后半)"
    total_col = header.index("总分") + 1
    assert ws.cell(row=rows - 2, column=total_col).value == 85.0  # (80+90)/2
    assert ws.cell(row=rows - 1, column=total_col).value == 90.0  # 前半
    assert ws.cell(row=rows, column=total_col).value == 80.0  # 后半

    # 页眉页脚
    assert "高一10班" in ws.oddHeader.left.text
    assert "测试" in ws.oddHeader.center.text
    assert ws.oddHeader.right.text == "2026年04月20日"
    assert "\n" in ws.oddFooter.left.text
    assert "班级平均分: " in ws.oddFooter.left.text
    assert "年级平均分: " in ws.oddFooter.left.text
    assert "\n 年级平均分: " in ws.oddFooter.left.text  # 第二行行首空格
    assert "(--|--)" in ws.oddFooter.left.text  # 无层次信息时兜底，无 A/B 标签
    assert "班级中位数: " in ws.oddFooter.right.text
    assert "(--|--) 年级中位数:" in ws.oddFooter.right.text
    assert "(标准差" not in ws.oddFooter.right.text  # 中位数不记录标准差
    assert "：" not in ws.oddFooter.left.text and "：" not in ws.oddFooter.right.text
    assert "（" not in ws.oddFooter.left.text and "（" not in ws.oddFooter.right.text

    # 字体与列宽
    assert ws.cell(row=1, column=1).font.name == "方正小标宋_GBK"
    assert ws.cell(row=1, column=1).font.size == 12
    assert ws.cell(row=2, column=1).font.name == "宋体"
    assert ws.cell(row=2, column=1).font.size == 11
    assert ws.cell(row=rows, column=1).font.name == "宋体"  # 平均行
    assert ws.column_dimensions["A"].width == 6
    assert ws.column_dimensions["B"].width == 8
    assert ws.column_dimensions["D"].width == 7
    assert ws.column_dimensions["F"].width == 5
    assert ws.column_dimensions["G"].width == 5
    assert ws.column_dimensions["K"].width == 5  # 校次
    assert ws.column_dimensions["L"].width == 5  # 班次
    assert ws.oddHeader.left.font == "微软雅黑,bold"
    assert ws.oddHeader.left.size == 20
    assert ws.oddHeader.center.font == "方正小标宋_GBK"
    assert ws.oddHeader.center.size == 20
    assert ws.oddHeader.right.size == 14
    assert ws.oddFooter.left.size == 12

    # 对齐与行高
    assert ws.row_dimensions[1].height == 20.0
    assert ws.cell(row=1, column=1).alignment.horizontal == "center"
    assert ws.cell(row=1, column=1).alignment.vertical == "center"
    assert ws.cell(row=2, column=3).alignment.horizontal == "center"  # 总分列
    assert ws.cell(row=rows, column=1).alignment.horizontal == "center"  # 平均行首格

    # 条件格式：客观分/主观分/单选/多选/小题/大题列数据条，范围到倒数第 4 行
    cfs = list(ws.conditional_formatting)
    assert len(cfs) == 7  # 客观/主观/单选/多选 + 26(1)/26(2)/26
    for cf in cfs:
        assert str(cf.sqref).endswith(str(rows - 3))
        for rule in cf.rules:
            assert "dataBar" in rule.type
            assert rule.dataBar.color.rgb == "0067C487"
    letters = {str(cf.sqref).split(":")[0][0] for cf in cfs}
    assert letters == {"D", "E", "F", "G", "H", "I", "J"}
    assert "C" not in letters  # 总分列不设

    # 表格框线：默认细实线；单选|多选、同大题小题之间取消竖线
    assert ws.cell(row=2, column=1).border.left.style == "thin"
    assert ws.cell(row=2, column=6).border.right is None  # F 右无竖线
    assert ws.cell(row=2, column=7).border.left is None  # G 左无竖线
    assert ws.cell(row=2, column=8).border.right is None  # 26(1) 右无竖线
    assert ws.cell(row=2, column=9).border.left is None  # 26(2) 左无竖线
    assert ws.cell(row=2, column=9).border.right.style == "thin"  # 26(2)|26 保留
