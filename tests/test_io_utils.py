"""底层读取工具的单元测试。"""

import pytest

from grade_analyzer.io_utils import read_raw_sheet, write_excel_report


def test_read_raw_sheet_real_joint_file():
    df = read_raw_sheet("data/input/测试样例/地理原始数据.xlsx")
    assert df.shape == (1364, 63)


def test_read_raw_sheet_real_weekly_file():
    df = read_raw_sheet(
        "data/input/测试样例/【教学班报告--高一下地理限时练一】所有班级学生小题得分明细.xlsx"
    )
    assert df.shape == (408, 45)


def test_read_raw_sheet_missing_file():
    with pytest.raises(FileNotFoundError):
        read_raw_sheet("data/input/不存在.xlsx")


def test_read_raw_sheet_bad_extension(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="扩展名"):
        read_raw_sheet(str(f))


def test_write_excel_report_multi_sheet(tmp_path):
    import pandas as pd

    path = tmp_path / "报告.xlsx"
    write_excel_report(
        {
            "科目统计": pd.DataFrame({"考试名称": ["期中"], "考生数": [3]}),
            "配置说明": pd.DataFrame({"配置项": ["及格线"], "值": [0.6]}),
        },
        str(path),
    )
    sheets = pd.read_excel(path, sheet_name=None)
    assert set(sheets) == {"科目统计", "配置说明"}


def test_write_excel_report_number_formats(tmp_path):
    import openpyxl
    import pandas as pd

    path = tmp_path / "格式.xlsx"
    df = pd.DataFrame({"平均分": [55.46], "及格率": [0.37995]})
    write_excel_report(
        {"统计": df},
        str(path),
        {"统计": {"平均分": "0.0", "及格率": "0.00%"}},
    )
    wb = openpyxl.load_workbook(path)
    ws = wb["统计"]
    header = {cell.value: cell.column for cell in ws[1] if cell.value}
    assert ws.cell(row=2, column=header["平均分"]).number_format == "0.0"
    assert ws.cell(row=2, column=header["及格率"]).number_format == "0.00%"
