"""客观题得分明细汇总与得分率距平分析核心实现。

输入：一个文件夹，内含各班 "N班.xls" 与 "全部班级.xls"（客观题得分明细）。
文件夹名形如 "<考试名>(<学科>)客观题得分明细" 或 "<考试名>客观题得分明细"，
考试规范名称按项目规则由 学期简称 + 学科 + 考试名 拼接（config.normalize_exam_name）。

输出：
- {考试规范名称}_客观题得分汇总（全部班级）.xlsx：每班一个 sheet + 得分率汇总；
- 每个教师一个 {考试规范名称}_客观题得分率距平（{教师}）.xlsx：班级对比 + 距平统计。

班级与教师映射来自全局学科配置（config/subjects/<学期>_<学科>.yaml）。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date as _date
from pathlib import Path

import pandas as pd
import xlrd
import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from grade_analyzer.cleaning import normalize_class_name
from grade_analyzer.config import load_config, normalize_exam_name
from grade_analyzer.detect import detect_subject_from_filename

FOLDER_SUFFIX = "客观题得分明细"
CLASS_SUM_NAME = "全部班级"
PCT = "0.00%"
NUM2 = "0.00"

_SIDE_MEDIUM = Side(style="medium", color="FF000000")
_SIDE_THIN = Side(style="thin", color="FF000000")
_TITLE_FONT = Font(name="微软雅黑", size=12, bold=True)
_SUBTITLE_FONT = Font(name="微软雅黑", size=11)
_BOLD_FONT = Font(name="微软雅黑", size=11, bold=True)
_CENTER = Alignment(vertical="center", horizontal="center")

_COLOR_SCALE_RULE = ColorScaleRule(
    start_type="percentile",
    start_value=10,
    start_color="FFFFFF",
    mid_type="percentile",
    mid_value=50,
    mid_color="CCCCCC",
    end_type="percentile",
    end_value=90,
    end_color="AAAAAA",
)
_DATA_BAR_RULE = DataBarRule(
    start_type="num",
    start_value=0,
    end_type="num",
    end_value=1,
    color="FF67C487",
)

_PLUGIN_CONFIG_KEYS = {
    "input_dir",
    "output_dir",
    "semester",
    "subject",
    "exam_date",
    "low_score_flag",
    "max_date_input_errors",
    "output_subdir_by_exam",
    "summary_groups",
    "deviation_groups",
    "baseline",
    "bl",
    "school",
}
_DATE_RE = re.compile(r"^(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?$")
_SUMMARY_ROWS = [
    "正距平题数",
    "正距平均值",
    "正距平总值",
    "负距平题数",
    "负距平均值",
    "负距平总值",
    "低负距平数",
    "低负距平表",
    "低得分率表",
    "低得分题数",
    "　均总得分",
    "总得分距平",
]


class ExamDateAborted(RuntimeError):
    """用户未确认使用今日日期，任务中止。"""


class InputDirAborted(RuntimeError):
    """输入目录错误次数过多，任务中止。"""


class SchoolInputAborted(RuntimeError):
    """学校选择错误次数过多，任务中止。"""


@dataclass
class ClassSheet:
    """一个班级文件解析结果。"""

    class_name: str
    header: list
    subheader: list
    rows: list


@dataclass
class GroupDef:
    """分组定义：类型 + 标签 + 成员班级短名。"""

    group_type: str  # teacher / level / unassigned / all
    name: str        # 柯 / A / 未分层 / 全部班级
    classes: list[str]


def _set_col(i: int) -> str:
    """0-based 列索引 -> Excel 列字母（支持超过 26 列）。"""
    if i < 26:
        return chr(ord("A") + i)
    return chr(ord("A") + i // 26 - 1) + chr(ord("A") + i % 26)


def trans_num(value):
    """把 "47.92%" / "0.96" 等文本转为数值，空串转 None，其余原样。"""
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return None
        if s.endswith("%"):
            try:
                return float(s[:-1]) / 100
            except ValueError:
                return value
        if re.fullmatch(r"-?\d+(\.\d+)?", s):
            try:
                return float(s)
            except ValueError:
                return value
    return value


def _split_folder_name(folder_name: str) -> tuple[str, str | None]:
    """按括号拆分文件夹名，返回 (考试名前缀, 括号内学科)。"""
    for open_ch, close_ch in (("（", "）"), ("(", ")")):
        if open_ch in folder_name and close_ch in folder_name:
            start = folder_name.index(open_ch)
            end = folder_name.index(close_ch)
            if start < end:
                prefix = folder_name[:start].strip()
                subject_raw = folder_name[start + 1 : end].strip()
                return prefix, subject_raw or None
    prefix = folder_name.strip()
    if prefix.endswith(FOLDER_SUFFIX):
        prefix = prefix[: -len(FOLDER_SUFFIX)].strip()
    return prefix, None


def _resolve_subject(raw: str | None, config) -> str | None:
    """把括号内学科映射到项目 subjects / 别名中的规范名称。"""
    if not raw:
        return None
    if raw in config.subjects:
        return raw
    for canonical, words in config.subject_aliases.items():
        if raw in words:
            return canonical
    raise ValueError(
        f"学科 {raw!r} 不在项目 subjects/别名配置中（可选: {'/'.join(config.subjects)}）"
    )


def derive_exam_info(
    folder_name: str,
    semester: str,
    config,
    subject_override: str | None = None,
) -> tuple[str, str | None]:
    """从文件夹名推导 (考试规范名称, 学科)。"""
    prefix, subject_raw = _split_folder_name(folder_name)
    subject = _resolve_subject(subject_raw, config)
    if subject is None:
        subject = _resolve_subject(subject_override, config)
    if not prefix:
        raise ValueError(f"无法从文件夹名 {folder_name!r} 提取考试名称")
    exam_name = normalize_exam_name(prefix, semester, subject, config.subject_aliases)
    if not exam_name:
        raise ValueError(f"考试规范名称解析失败：{folder_name!r}")
    return exam_name, subject


def _short_class_name(class_name: str) -> str:
    """高一01班 -> 1班；已是短名则原样（去掉班级号前导零以匹配文件名）。"""
    m = re.search(r"\d+", class_name)
    return f"{int(m.group(0))}班" if m else class_name


def _parse_group_types(value, label: str) -> list[str]:
    """解析分组类型配置（teacher/level 可多选）。"""
    if value is None:
        return []
    if isinstance(value, str):
        value = [v.strip() for v in value.split(",") if v.strip()]
    if not isinstance(value, list):
        raise ValueError(f"{label} 应为列表（teacher/level）")
    result: list[str] = []
    for v in value:
        v = str(v).strip()
        if v not in ("teacher", "level"):
            raise ValueError(f"{label} 仅支持 teacher/level，当前为 {v!r}")
        if v not in result:
            result.append(v)
    return result


def build_groups(
    config,
    semester: str,
    subject: str | None,
    class_names: list[str],
    group_types: list[str],
) -> dict[str, list[GroupDef]]:
    """按分组类型把输入班级拆分为分组（全部班级为特殊分组，不在此列）。"""
    result: dict[str, list[GroupDef]] = {}
    if "teacher" in group_types and subject:
        teacher_map = config.teacher_maps.get((semester, subject))
        if teacher_map is None:
            print(
                f"[客观题] 未找到学科配置 config/subjects/{semester}_{subject}.yaml，"
                "跳过教师分组"
            )
        else:
            for code, name in teacher_map.teacher_names.items():
                classes = [
                    _short_class_name(cls)
                    for cls, c in teacher_map.class_teachers.items()
                    if c == code
                ]
                classes = [c for c in classes if c in class_names]
                if classes:
                    result.setdefault("teacher", []).append(
                        GroupDef("teacher", name, classes)
                    )
    if "level" in group_types:
        level_map: dict[str, list[str]] = {}
        for (sem, full), info in config.class_infos.items():
            if sem != semester:
                continue
            short = _short_class_name(full)
            if short in class_names:
                level_map.setdefault(info.level, []).append(short)
        for level in sorted(level_map):
            result.setdefault("level", []).append(
                GroupDef("level", level, level_map[level])
            )
        assigned = {c for v in level_map.values() for c in v}
        unassigned = [c for c in class_names if c not in assigned]
        if unassigned:
            result.setdefault("level", []).append(
                GroupDef("unassigned", "未分层", unassigned)
            )
    return result


def resolve_baseline(value: str, groups_by_type: dict[str, list[GroupDef]]) -> GroupDef:
    """把基线配置值解析为分组（全部班级 或 分组名）。"""
    v = (value or "").strip()
    if not v or v == CLASS_SUM_NAME:
        return GroupDef("all", CLASS_SUM_NAME, [])
    for groups in groups_by_type.values():
        for group in groups:
            if group.name == v:
                return group
    available = "、".join(
        sorted({g.name for groups in groups_by_type.values() for g in groups})
    )
    raise ValueError(f"基线 {v!r} 无法解析（可用基线: {available}、{CLASS_SUM_NAME}）")


def teachers_from_config(config, semester: str, subject: str | None) -> dict[str, list[str]]:
    """从全局学科配置读取 教师名 -> 班级短名列表。"""
    if not subject:
        return {}
    teacher_map = config.teacher_maps.get((semester, subject))
    if teacher_map is None:
        raise ValueError(f"未找到学科配置 config/subjects/{semester}_{subject}.yaml")
    teachers: dict[str, list[str]] = {}
    for code, name in teacher_map.teacher_names.items():
        classes = [
            _short_class_name(cls)
            for cls, c in teacher_map.class_teachers.items()
            if c == code
        ]
        if classes:
            teachers[name] = classes
    return teachers


def _class_sort_key(path: Path) -> tuple[int, int]:
    if path.stem == CLASS_SUM_NAME:
        return (1, 0)
    m = re.fullmatch(r"(\d+)班", path.stem)
    return (0, int(m.group(1))) if m else (2, 0)


def discover_class_files(input_dir: Path) -> list[Path]:
    """扫描输入文件夹，按数值排序返回 N班.xls 与 全部班级.xls。"""
    if not input_dir.is_dir():
        raise FileNotFoundError(f"输入文件夹不存在: {input_dir}")
    files = [
        p
        for p in input_dir.iterdir()
        if p.is_file()
        and p.suffix.lower() == ".xls"
        and (p.stem == CLASS_SUM_NAME or re.fullmatch(r"\d+班", p.stem))
    ]
    if not files:
        raise ValueError(f"输入文件夹中未找到 N班.xls / 全部班级.xls 文件: {input_dir}")
    return sorted(files, key=_class_sort_key)


def read_class_sheet(path: Path) -> ClassSheet:
    """xlrd 读取单个班级 .xls（前 2 行跳过，第 3 行为表头，第 4 行为子表头）。"""
    wb = xlrd.open_workbook(str(path))
    ws = wb.sheet_by_index(0)
    if ws.nrows < 5:
        raise ValueError(f"{path}: 数据行数异常（{ws.nrows} 行）")
    ncols = ws.ncols
    header = [trans_num(ws.cell_value(2, c)) for c in range(ncols)]
    subheader = [trans_num(ws.cell_value(3, c)) for c in range(ncols)]
    rows = [
        [trans_num(ws.cell_value(r, c)) for c in range(ncols)]
        for r in range(4, ws.nrows)
    ]
    return ClassSheet(
        class_name=path.stem, header=header, subheader=subheader, rows=rows
    )


def _set_page_print(ws, exam_name: str, exam_date: str, left: str = "") -> None:
    """设置页眉、横向打印、打印区域与缩放（复刻参考输出）。"""
    ws.print_options.horizontalCentered = True
    if left:
        ws.oddHeader.left.text = left
        ws.oddHeader.left.size = 18
        ws.oddHeader.left.font = "微软雅黑,bold"
    ws.oddHeader.center.text = exam_name
    ws.oddHeader.center.size = 20
    ws.oddHeader.center.font = "方正小标宋_GBK"
    if exam_date:
        ws.oddHeader.right.text = exam_date
        ws.oddHeader.right.size = 13
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = True
    ws.print_area = f"A1:{_set_col(ws.max_column - 1)}{ws.max_row}"
    ws.sheet_view.zoomScale = 150


def _style_class_sheet(
    ws, n_questions: int, class_name: str, exam_name: str, exam_date: str
) -> None:
    """复刻参考输出的每班 sheet 样式。"""
    ws.freeze_panes = "B3"
    for i in range(6):
        cell_col = _set_col(i)
        ws.merge_cells(f"{cell_col}1:{cell_col}2")
        ws[f"{cell_col}1"].alignment = _CENTER
        ws[f"{cell_col}1"].font = _TITLE_FONT
    ws.merge_cells("G1:L1")
    ws["G1"].alignment = _CENTER
    ws["G1"].font = _TITLE_FONT
    for col in "GHIJKL":
        ws[f"{col}2"].alignment = _CENTER
        ws[f"{col}2"].font = _SUBTITLE_FONT
        if col in "GHIJ":
            ws.column_dimensions[col].width = 12

    last_data_row = n_questions + 2
    for r in range(2, last_data_row + 1):
        for col in "DGHIJKL":
            ws[f"{col}{r}"].number_format = PCT
        ws[f"E{r}"].number_format = NUM2
        ws.conditional_formatting.add(f"G{r}:J{r}", _DATA_BAR_RULE)
    for r in range(3, last_data_row + 1):
        answer = ws.cell(r, 6).value
        for ch in str(answer):
            if "A" <= ch <= "D":
                ws[f"{_set_col(ord(ch) - ord('A') + 6)}{r}"].font = _BOLD_FONT
    ws.conditional_formatting.add(f"D3:D{last_data_row}", _COLOR_SCALE_RULE)

    max_row, max_col = ws.max_row, ws.max_column
    for j in range(1, max_row + 1):
        ws.row_dimensions[j].height = 16.5
        for i in range(max_col):
            cell = ws[f"{_set_col(i)}{j}"]
            top = _SIDE_MEDIUM if j == 1 else _SIDE_THIN
            cell.border = Border(top=top, bottom=_SIDE_THIN, left=None, right=None)
            if i not in (3, 6, 7, 8, 9):
                cell.alignment = Alignment(horizontal="center")
    _set_page_print(ws, exam_name, exam_date, left=class_name)


def _build_score_frame(sheets: list[ClassSheet]) -> pd.DataFrame:
    """题号 x 班级 得分率透视表，末尾追加 总分（平均分之和）行。"""
    score_list: dict = {}
    ave_lst: dict = {}
    for i, sheet in enumerate(sheets):
        if i == 0:
            score_list["题号"] = [row[0] for row in sheet.rows]
            score_list["题型"] = [row[1] for row in sheet.rows]
        score_list[sheet.class_name] = [row[3] for row in sheet.rows]
        ave_lst[sheet.class_name] = sum(
            v
            for v in (row[4] for row in sheet.rows)
            if isinstance(v, (int, float)) and v == v
        )
    score_df = pd.DataFrame(score_list).set_index("题号")
    score_df.loc["总分", :] = ave_lst
    return score_df


def _write_summary_sheet(wb: Workbook, summary_df: pd.DataFrame, exam_name: str, exam_date: str) -> None:
    """写入并样式化 得分率汇总 sheet。"""
    ws = wb.create_sheet(title="得分率汇总")
    for row in dataframe_to_rows(summary_df, index=True):
        if len(row) < 2:
            continue
        ws.append(row)
    n_questions = len(summary_df) - 1
    ws.column_dimensions["A"].width = 6.5
    ws.column_dimensions["B"].width = 6.5
    for i in range(2, len(summary_df.columns) + 1):
        cell_col = _set_col(i)
        ws.column_dimensions[cell_col].width = 7.8
        ws.conditional_formatting.add(
            f"{cell_col}2:{cell_col}{n_questions + 1}", _COLOR_SCALE_RULE
        )
        for r in range(2, n_questions + 2):
            ws[f"{cell_col}{r}"].number_format = PCT
    all_col = _set_col(len(summary_df.columns))
    last_class_col = _set_col(len(summary_df.columns) - 1)
    max_row, max_col = ws.max_row, ws.max_column
    for j in range(max_row):
        ws.row_dimensions[j + 1].height = 20
        top = _SIDE_MEDIUM if j == 0 else _SIDE_THIN
        for i in range(max_col):
            col = _set_col(i)
            left = _SIDE_THIN if col == all_col else None
            right = _SIDE_THIN if col == last_class_col else None
            ws[f"{col}{j + 1}"].border = Border(
                top=top, bottom=_SIDE_THIN, left=left, right=right
            )
    for i in range(len(summary_df.columns) + 1):
        cell_col = _set_col(i)
        ws[f"{cell_col}1"].font = _TITLE_FONT
        ws[f"{cell_col}1"].alignment = _CENTER
    _set_page_print(ws, exam_name, exam_date)


def build_summary_workbook(
    unit_sheets: list[ClassSheet], exam_name: str, exam_date: str
) -> Workbook:
    """生成一个汇总工作簿：成员班级 sheet + 全部班级 sheet + 得分率汇总。"""
    wb = Workbook()
    for i, sheet in enumerate(unit_sheets):
        ws = wb.active if i == 0 else wb.create_sheet(title=sheet.class_name)
        if i == 0:
            ws.title = sheet.class_name
        ws.append(sheet.header)
        ws.append(sheet.subheader)
        for row in sheet.rows:
            ws.append(row)
        _style_class_sheet(ws, len(sheet.rows), sheet.class_name, exam_name, exam_date)
    _write_summary_sheet(
        wb, _build_score_frame(unit_sheets), exam_name, exam_date
    )
    return wb


def _build_summary_units(
    sheets: list[ClassSheet],
    groups_by_type: dict[str, list[GroupDef]],
    summary_types: list[str],
) -> list[tuple[str, list[ClassSheet]]]:
    """汇总输出单元：无分组 → [全部班级]（全部班级 sheet 始终包含）；分组 → 每组一个。"""
    if not summary_types:
        return [(CLASS_SUM_NAME, sheets)]
    all_sheet = [s for s in sheets if s.class_name == CLASS_SUM_NAME]
    units: list[tuple[str, list[ClassSheet]]] = []
    for t in summary_types:
        for group in groups_by_type.get(t, []):
            member = [s for s in sheets if s.class_name in group.classes]
            if member:
                units.append((group.name, member + all_sheet))
    return units


def _deviation_summary(
    clses_df: pd.DataFrame,
    class_names: list[str],
    baseline_col: str,
    low_score_flag: float,
) -> tuple[pd.DataFrame, int, int]:
    """计算教师班级距平统计表，返回 (汇总表, 最大低负距平数, 最大低得分题数)。"""
    data = clses_df.iloc[:-1]  # 统计不含总分行
    summary: dict[str, list] = {}
    max_low_low_count = 0
    max_low_score_count = 0
    for cls in class_names:
        dev = data[f"{cls}_d"]
        low = dev[dev < 0]
        low_count = int(low.count())
        low_ave = float(low.mean()) if low_count else float("nan")
        low_sum = float(low.sum()) if low_count else 0.0
        up = dev[dev > 0]
        up_count = int(up.count())
        up_ave = float(up.mean()) if up_count else float("nan")
        up_sum = float(up.sum()) if up_count else 0.0

        orig = data[cls]
        low_score = orig[orig <= low_score_flag]
        low_low = low[low < low_ave]
        low_low1 = low_low[[i for i in low_low.index if i not in low_score.index]]
        low_low2 = low_low[[i for i in low_low.index if i in low_score.index]]
        low_score3 = low_score[[i for i in low_score.index if i in low_low.index]]
        low_score4 = low_score[[i for i in low_score.index if i not in low_low.index]]
        low_low_map = pd.concat((low_low1, low_low2)).to_dict()
        low_low_lst = [f"{k: >6}({v: >7.2%})" for k, v in low_low_map.items()]
        max_low_low_count = max(max_low_low_count, len(low_low_map))
        low_score_map = pd.concat((low_score3, low_score4)).to_dict()
        low_score_lst = [f"{k: >6} ({v: >6.2%})" for k, v in low_score_map.items()]
        max_low_score_count = max(max_low_score_count, len(low_score_lst))

        sum_score = float(clses_df.loc["总分", cls])
        sum_score_m = sum_score - float(clses_df.loc["总分", baseline_col])
        summary[f"{cls}_d"] = [
            up_count,
            up_ave,
            up_sum,
            low_count,
            low_ave,
            low_sum,
            len(low_low_map),
            "\n".join(low_low_lst),
            "\n".join(low_score_lst),
            len(low_score_lst),
            sum_score,
            sum_score_m,
        ]
    return pd.DataFrame(summary, index=_SUMMARY_ROWS), max_low_low_count, max_low_score_count


def _style_deviation_sheet(ws, cls_count: int, max_row: int) -> None:
    """复刻参考输出文件 2 的班级对比 sheet 样式。"""
    for i in range(cls_count * 2 + 2):
        cell_col = _set_col(i)
        ws[f"{cell_col}1"].alignment = Alignment(horizontal="center")
        ws[f"{cell_col}1"].font = Font(bold=True)
        if i >= cls_count + 2:
            ws.conditional_formatting.add(
                f"{cell_col}2:{cell_col}{max_row}", _COLOR_SCALE_RULE
            )
        for j in range(1, max_row):
            ws[f"{cell_col}{j + 1}"].number_format = PCT
    ws.sheet_view.zoomScale = 150


def _style_deviation_summary(
    ws,
    cls_count: int,
    max_low_low_count: int,
    max_low_score_count: int,
    exam_name: str,
    exam_date: str,
) -> None:
    """复刻参考输出 距平统计 sheet 样式。"""
    ws.column_dimensions["A"].width = 12
    ws.row_dimensions[9].height = 14 * max_low_low_count
    ws.row_dimensions[10].height = 14 * max_low_score_count
    for i in range(cls_count + 1):
        col = _set_col(i)
        for row in range(1, len(_SUMMARY_ROWS) + 2):
            cell = f"{col}{row}"
            top = _SIDE_MEDIUM if row == 1 else _SIDE_THIN
            right = _SIDE_THIN if col == "A" else None
            ws[cell].border = Border(top=top, bottom=_SIDE_THIN, left=None, right=right)
            if col == "A":
                continue
            ws[f"{col}9"].alignment = Alignment(vertical="top", wrapText=True)
            ws[f"{col}10"].alignment = Alignment(vertical="top", wrapText=True)
            ws[f"{col}1"].alignment = Alignment(horizontal="center")
            ws[f"{col}1"].font = Font(bold=True)
            if row in (3, 4, 6, 7):
                ws.column_dimensions[col].width = 18
                ws[cell].number_format = PCT
    ws["A9"].alignment = Alignment(vertical="center")
    ws["A10"].alignment = Alignment(vertical="center")
    ws.print_options.horizontalCentered = True
    ws.oddHeader.center.text = exam_name
    ws.oddHeader.center.size = 20
    ws.oddHeader.center.font = "方正小标宋_GBK"
    ws.oddHeader.right.text = exam_date
    ws.oddHeader.right.size = 13
    ws.sheet_view.zoomScale = 150


def build_deviation_workbook(
    score_df: pd.DataFrame,
    group: GroupDef,
    baseline: GroupDef,
    low_score_flag: float,
    exam_name: str,
    exam_date: str,
) -> Workbook:
    """生成输出文件 2：分组班级得分率对比（基线=全部班级或指定分组）+ 距平统计。"""
    class_names = group.classes
    if baseline.group_type == "all":
        baseline_col = CLASS_SUM_NAME
        clses_df = score_df.loc[:, [*class_names, baseline_col]].copy()
    else:
        baseline_col = baseline.name
        clses_df = score_df[class_names].copy()
        cols = [c for c in baseline.classes if c in score_df.columns]
        if not cols:
            raise ValueError(f"基线分组 {baseline.name!r} 无可用班级列")
        clses_df[baseline_col] = score_df[cols].mean(axis=1)
    for cls in class_names:
        clses_df[f"{cls}_d"] = clses_df[cls] - clses_df[baseline_col]

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    for row in dataframe_to_rows(clses_df, index=True):
        if len(row) < 2:
            continue
        ws.append(row)
    _style_deviation_sheet(ws, len(class_names), len(clses_df))

    summary_df, max_low_low, max_low_score = _deviation_summary(
        clses_df, class_names, baseline_col, low_score_flag
    )
    ws2 = wb.create_sheet(title="距平统计")
    for row in dataframe_to_rows(summary_df, index=True):
        if len(row) < 2:
            continue
        ws2.append(row)
    _style_deviation_summary(
        ws2,
        len(class_names),
        max_low_low,
        max_low_score,
        exam_name,
        exam_date,
    )
    return wb


def _default_plugin_config_path() -> str:
    """插件自带配置文件路径（plugins/objective_analyze/config.yaml）。"""
    return str(Path(__file__).resolve().parent / "config.yaml")


def _load_plugin_config(path: str | None) -> dict:
    """读取插件配置；未指定路径时自动读取插件目录下的 config.yaml（存在才读取）。"""
    if not path:
        path = _default_plugin_config_path()
        if not Path(path).is_file():
            return {}
    cfg_path = Path(path)
    if not cfg_path.is_file():
        raise FileNotFoundError(f"插件配置文件不存在: {cfg_path}")
    with open(cfg_path, encoding="utf-8") as fh:
        raw = yaml.safe_load(fh) or {}
    if not isinstance(raw, dict):
        raise ValueError(f"{cfg_path}: 插件配置应为映射")
    unknown = set(raw) - _PLUGIN_CONFIG_KEYS
    if unknown:
        raise ValueError(f"{cfg_path}: 未知字段 {sorted(unknown)}")
    return raw


def _format_exam_date(value: str | None) -> str:
    if not value:
        value = _date.today().isoformat()
    try:
        return _date.fromisoformat(str(value)).strftime("%Y年%m月%d日")
    except ValueError:
        raise ValueError(f"exam_date 应为 YYYY-MM-DD，当前为 {value!r}")


def _parse_date(value) -> str:
    """解析并归一化日期输入（如 2026/5/20、2026年5月20日 -> 2026-05-20）。"""
    m = _DATE_RE.match(str(value).strip())
    if not m:
        raise ValueError(
            f"日期格式不正确: {value!r}（应为 YYYY-MM-DD 或 2026年5月20日 等）"
        )
    try:
        return _date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
    except ValueError:
        raise ValueError(f"日期不存在: {value!r}")


def find_exam_date(config, exam_name: str, semester: str, subject: str | None, prefix: str) -> str | None:
    """从考试配置目录（exams_dir）搜索匹配条目的日期。

    匹配优先级：
    1. 条目规范名称 == 推导的考试规范名称；
    2. 学期 + 学科一致，且 short_name 等于文件夹前缀（或名称包含前缀）。
    无匹配返回 None；优先级 2 多个匹配时报错。
    """
    from grade_analyzer.config import _load_exam_file, discover_exam_files

    candidates: list[tuple[str, str]] = []
    for file_path, folder_semester in discover_exam_files(config.exams_dir):
        try:
            entry = _load_exam_file(
                file_path, folder_semester, config.subject_aliases, config.input_dir
            )
        except (ValueError, FileNotFoundError) as exc:
            print(f"[客观题] 跳过无法解析的考试条目 {file_path}: {exc}")
            continue
        if not entry.date:
            continue
        if entry.name == exam_name:
            return entry.date
        if (
            entry.semester == semester
            and entry.subject == subject
            and (
                entry.short_name == prefix
                or (entry.name and prefix and prefix in entry.name)
            )
        ):
            candidates.append((str(file_path), entry.date))
    if len(candidates) == 1:
        return candidates[0][1]
    if len(candidates) > 1:
        choices = "；".join(f"{path}（{date}）" for path, date in candidates)
        raise ValueError(f"考试配置目录中存在多个匹配条目，无法确定考试日期: {choices}")
    return None


def _ask_exam_date(max_errors: int = 2) -> str:
    """交互输入考试日期；错误达上限后询问是否用今日日期，未确认则中止。"""
    errors = 0
    while True:
        try:
            value = input("请输入考试日期（如 2026-05-20）：").strip().lstrip("\ufeff")
        except EOFError:
            print("[客观题] 未获取到输入，已中止")
            raise ExamDateAborted()
        try:
            return _parse_date(value)
        except ValueError as exc:
            errors += 1
            print(f"[客观题] {exc}")
            if errors >= max_errors:
                try:
                    choice = input("是否使用今日日期作为考试日期？（y/n）：").strip().lower()
                except EOFError:
                    print("[客观题] 未获取到输入，已中止")
                    raise ExamDateAborted()
                if choice in ("y", "yes", "是"):
                    return _date.today().isoformat()
                raise ExamDateAborted()


def _ask_input_dir(max_errors: int = 2) -> str:
    """交互输入数据源文件夹；错误达上限后中止。"""
    errors = 0
    while True:
        try:
            value = input("请输入数据源文件夹路径：").strip().lstrip("\ufeff")
        except EOFError:
            print("[客观题] 未获取到输入，已中止")
            raise InputDirAborted()
        if not value:
            errors += 1
            print("[客观题] 输入不能为空")
            if errors >= max_errors:
                raise InputDirAborted()
            continue
        path = Path(value)
        if path.is_dir():
            return str(path)
        errors += 1
        print(f"[客观题] 文件夹不存在或不是目录: {value!r}")
        if errors >= max_errors:
            raise InputDirAborted()


_RAW_EXTS = {".xls", ".xlsx", ".csv"}


def _resolve_input(
    input_dir_arg: str | None, raw: dict, max_errors: int
) -> tuple[Path, str]:
    """解析输入源，返回 (路径, mode)。

    mode: class_files=分班明细文件夹；detail=小题分（含答案）原始文件。
    兼容：文件→detail；文件夹含 N班.xls→class_files；
    文件夹无 N班.xls→交互提示输入原始成绩文件名称→detail。
    """
    src = (input_dir_arg or raw.get("input_dir") or "").strip()
    if src:
        path = Path(src)
        if path.is_file():
            if path.suffix.lower() not in _RAW_EXTS:
                raise ValueError(
                    f"不支持的原始成绩文件类型: {path.suffix}（支持 {sorted(_RAW_EXTS)}）"
                )
            return path, "detail"
        if path.is_dir():
            try:
                class_files = discover_class_files(path)
            except ValueError:
                class_files = []
            if class_files:
                return path, "class_files"
            print(
                "[客观题] 文件夹中未发现 N班.xls / 全部班级.xls，"
                "按小题分（含答案）原始文件流程处理"
            )
            return _ask_raw_file_name(path, max_errors=max_errors), "detail"
        raise FileNotFoundError(f"输入路径不存在: {path}")
    print("[客观题] 未指定数据源，请输入数据源文件夹或小题分原始文件路径")
    return _ask_source(max_errors=max_errors)


def _ask_source(max_errors: int) -> tuple[Path, str]:
    """交互输入数据源（文件夹或小题分原始文件），错误达上限后中止。"""
    errors = 0
    while True:
        try:
            value = input("请输入数据源文件夹或小题分原始文件路径：").strip()
        except EOFError:
            raise InputDirAborted()
        if not value:
            errors += 1
            print("[客观题] 输入不能为空")
            if errors >= max_errors:
                raise InputDirAborted()
            continue
        path = Path(value)
        if path.is_file() and path.suffix.lower() in _RAW_EXTS:
            return path, "detail"
        if path.is_dir():
            try:
                class_files = discover_class_files(path)
            except ValueError:
                class_files = []
            if class_files:
                return path, "class_files"
            return _ask_raw_file_name(path, max_errors=max_errors), "detail"
        errors += 1
        print(f"[客观题] 路径不存在或不是有效数据源: {value!r}")
        if errors >= max_errors:
            raise InputDirAborted()


def _ask_raw_file_name(folder: Path, max_errors: int = 2) -> Path:
    """在给定文件夹中交互输入小题分原始文件名称（兼容子目录相对路径）。"""
    errors = 0
    while True:
        try:
            value = input(f"请输入原始成绩文件名称（位于 {folder}，可含子目录相对路径）：").strip()
        except EOFError:
            raise InputDirAborted()
        if not value:
            errors += 1
            print("[客观题] 输入不能为空")
            if errors >= max_errors:
                raise InputDirAborted()
            continue
        path = folder / value
        if path.is_file() and path.suffix.lower() in _RAW_EXTS:
            return path
        errors += 1
        print(f"[客观题] 文件不存在或类型不支持: {path}")
        if errors >= max_errors:
            raise InputDirAborted()


def _parse_bool(value, default: bool = True) -> bool:
    if value is None:
        return default
    if isinstance(value, str):
        return value.strip().lower() not in ("", "false", "0", "no", "否")
    return bool(value)


def _resolve_output_dir(output_dir_arg: str | None, raw: dict, exam_name: str) -> str:
    """解析输出目录：--output-dir > 运行时确认（默认/回退 = 配置项或标准默认）。

    output_subdir_by_exam=True（默认）时，最终目录 = 基础目录/<考试规范名称>/。
    """
    subdir_by_exam = _parse_bool(raw.get("output_subdir_by_exam"), True)
    if output_dir_arg:
        base = str(Path(output_dir_arg))
        return str(Path(base) / exam_name) if subdir_by_exam else base
    else:
        configured = (raw.get("output_dir") or "").strip()
        base = configured or "data/output/task"
    default = str(Path(base) / exam_name) if subdir_by_exam else base
    try:
        value = input(f"请输入输出目录（回车使用默认 {default}）：").strip().lstrip("\ufeff")
    except EOFError:
        print(f"[客观题] 未获取到输入，使用默认输出目录: {default}")
        return default
    if not value:
        return default
    try:
        path = Path(value)
        path.mkdir(parents=True, exist_ok=True)
    except (OSError, ValueError):
        print(f"[客观题] 输出目录无效: {value!r}，使用默认输出目录: {default}")
        return default
    if subdir_by_exam:
        return str(Path(value) / exam_name)
    return str(path)


def _resolve_exam_date(
    raw: dict,
    config,
    exam_name: str,
    semester: str,
    subject: str | None,
    prefix: str,
) -> str:
    """按优先级解析考试日期：插件配置显式 > 考试配置目录 > 交互输入。"""
    explicit = raw.get("exam_date")
    if explicit:
        return _parse_date(explicit)
    found = find_exam_date(config, exam_name, semester, subject, prefix)
    if found:
        print(f"[客观题] 从考试配置目录匹配到考试日期: {found}")
        return found
    print("[客观题] 考试配置目录中未找到匹配条目，请输入考试日期")
    return _ask_exam_date(max_errors=int(raw.get("max_date_input_errors", 2)))



def _read_detail_frame(path: Path) -> pd.DataFrame:
    """读取小题分原始文件（xls/xlsx/csv，不处理表头）。"""
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path, header=None, dtype=object)
    return pd.read_excel(path, header=None, dtype=object, sheet_name=0)


def _detail_header_idx(raw: pd.DataFrame) -> int | None:
    """定位小题分表头行（含 姓名/考号）。"""
    for i, row in raw.iterrows():
        cells = {str(v).strip() for v in row.tolist()}
        if {"姓名", "考号"} <= cells:
            return i
    return None


def _is_number(value) -> bool:
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


def _normalize_answer(value) -> str:
    """答案归一：去空白、大写；仅保留 A-D 字母组合。"""
    if value is None:
        return ""
    text = re.sub(r"\s+", "", str(value)).upper()
    return text if re.fullmatch(r"[A-D]+", text) else ""


def _match_detail_exam(config, semester: str, subject: str | None, file_name: str):
    """在考试配置目录中匹配小题分文件对应条目：文件名优先，其次学期+科目唯一。"""
    base = Path(file_name).name
    candidates = []
    for exam in config.exams:
        if exam.semester != semester:
            continue
        if exam.subject and subject and exam.subject != subject:
            continue
        if Path(exam.file or "").name == base:
            return exam
        candidates.append(exam)
    if len(candidates) == 1:
        return candidates[0]
    return None


def _detail_fallback_name(path: Path, semester: str, subject: str, config) -> str | None:
    """从文件名推导考试名称（去常见后缀后按规范拼接）；失败返回 None。"""
    stem = path.stem
    for suffix in ("原始数据", "成绩数据", "小题得分明细", "成绩"):
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
            break
    stem = stem.strip("-_ ")
    if not stem:
        return None
    return normalize_exam_name(stem, semester, subject, config.subject_aliases)



def _distinct_schools(path: Path) -> list[str]:
    """返回小题分文件中出现过的学校列表（无学校列返回空）。"""
    raw = _read_detail_frame(path)
    hidx = _detail_header_idx(raw)
    if hidx is None:
        return []
    header = [str(v).strip() for v in raw.iloc[hidx].tolist()]
    if "学校" not in header:
        return []
    school_col = header.index("学校")
    schools = set()
    for v in raw.iloc[hidx + 1 :, school_col].tolist():
        text = str(v).strip() if v is not None else ""
        if text and text not in ("nan", "None"):
            schools.add(text)
    return sorted(schools)


def _validate_detail_school(path: Path, school: str) -> None:
    """校验给定学校存在于小题分数据中。"""
    schools = _distinct_schools(path)
    if schools and school not in schools:
        raise ValueError(f"学校 {school!r} 不在数据中（可选: {'、'.join(schools)}）")


def _resolve_detail_school(path: Path, config, max_errors: int) -> str | None:
    """交互选择筛选学校（存在性校验，失败重输，错误达上限后中止）。

    数据无学校列时返回 None（不过滤）。默认值：config.default_school 若在数据中，
    否则取第一个学校。
    """
    schools = _distinct_schools(path)
    if not schools:
        return None
    default = (config.default_school or "").strip()
    if default not in schools:
        default = schools[0]
    print(f"[客观题] 数据中的学校: {'、'.join(schools)}")
    errors = 0
    while True:
        try:
            value = input(f"请输入用于筛选的学校（回车使用 {default}）：").strip()
        except EOFError:
            raise SchoolInputAborted()
        if not value:
            value = default
        if value in schools:
            return value
        errors += 1
        print(f"[客观题] 学校 {value!r} 不存在，请重新输入")
        if errors >= max_errors:
            raise SchoolInputAborted()



def build_detail_class_sheets(
    path: Path,
    config,
    semester: str,
    subject: str | None,
    school: str | None = None,
) -> list:
    """从小题分（含小题答案）原始文件生成与分班文件同构的班级明细。

    规则：每题满分=该题全体最高得分；得分率=平均分/满分（缺考不计）；
    正确答案按班取满分数作答中的唯一答案，无法唯一判定时留空；
    题型按答案长度推断（多字母→多选题）；选项分布按 A/B/C/D/多选/未选。
    """
    raw = _read_detail_frame(path)
    hidx = _detail_header_idx(raw)
    if hidx is None:
        raise ValueError(f"{path}: 未找到表头行（需要 姓名/考号 列）")
    header = [str(v).strip() for v in raw.iloc[hidx].tolist()]
    col = {name: j for j, name in enumerate(header)}
    for need in ("姓名", "考号", "班级"):
        if need not in col:
            raise ValueError(f"{path}: 表头缺少 {need} 列")
    data = raw.iloc[hidx + 1 :].copy()
    id_col = col["考号"]
    data = data[data.iloc[:, id_col].map(
        lambda v: not (v is None or str(v).strip() in ("", "nan", "None"))
    )].reset_index(drop=True)
    if data.empty:
        raise ValueError(f"{path}: 未找到数据行")
    class_col = col["班级"]
    filter_school = school or (config.default_school if "学校" in col else None)
    if "学校" in col and filter_school:
        sch = data.iloc[:, col["学校"]].map(
            lambda v: str(v).strip() if v is not None else ""
        )
        matched = data[sch == filter_school]
        if matched.empty:
            raise ValueError(f"学校 {filter_school} 在数据中无记录")
        data = matched.reset_index(drop=True)
    candidates = []
    for j, h in enumerate(header):
        h2 = re.sub(r"\.0$", "", h) if re.fullmatch(r"\d+\.0", h) else h
        if re.fullmatch(r"\d+", h2):
            candidates.append((j, h2))
    def _kind(values, want):
        vals = [str(v).strip() for v in values
                if v is not None and str(v).strip() not in ("", "nan", "None")]
        if not vals:
            return False
        if want == "letters":
            return all(re.fullmatch(r"[A-Da-d]+", v) for v in vals)
        return all(_is_number(v) for v in vals)
    ans_map: dict[str, int] = {}
    score_map: dict[str, int] = {}
    for j, h2 in candidates:
        sample = data.iloc[:15, j].tolist()
        if _kind(sample, "letters"):
            ans_map.setdefault(h2, j)
        elif _kind(sample, "scores"):
            score_map.setdefault(h2, j)
    if not score_map:
        raise ValueError(f"{path}: 未找到客观题得分列（纯数字题号列）")
    if not ans_map:
        print("[客观题] 警告: 未识别到答案列，正确答案/选项分布将置空")
    question_labels = sorted(set(score_map) | set(ans_map), key=lambda x: int(x))
    def _norm_class(v) -> str:
        text = str(v).strip() if v is not None else ""
        cls, _g = normalize_class_name(text, config.default_grade)
        return cls
    data["_cls"] = data.iloc[:, class_col].map(_norm_class)
    class_names = sorted({c for c in data["_cls"].tolist() if c},
                         key=lambda c: (int(re.search(r"\d+", c).group()), c))
    header_row = ["题号", "题型", "分值", "得分率", "平均分", "正确答案",
                 "选项分布", "", "", "", "", ""]
    subheader_row = ["", "", "", "", "", "", "A", "B", "C", "D", "多选", "未选"]
    qmeta = []
    for label in question_labels:
        scol = score_map[label]
        acol = ans_map.get(label)
        full = float(pd.to_numeric(data.iloc[:, scol], errors="coerce").max())
        qmeta.append((int(label), full, scol, acol))
    sheets = []
    used_names = set()
    for cls in class_names + [CLASS_SUM_NAME]:
        if cls == CLASS_SUM_NAME:
            sheet_name = CLASS_SUM_NAME
        else:
            short = _short_class_name(cls)
            if short in used_names:
                print(f"[客观题] 警告: 班级 {cls} 与已有班级短名重复，使用完整班级名")
                sheet_name = cls
            else:
                sheet_name = short
        used_names.add(sheet_name)
        subset = data if cls == CLASS_SUM_NAME else data[data["_cls"] == cls]
        rows = []
        for label, full, scol, acol in qmeta:
            scores = pd.to_numeric(subset.iloc[:, scol], errors="coerce")
            mean = float(scores.mean()) if scores.notna().any() else 0.0
            rate = mean / full if full > 0 else 0.0
            ans = None
            if acol is not None:
                ans = data.iloc[:, acol].map(_normalize_answer).reindex(subset.index).fillna("")
            correct = ""
            if ans is not None and full > 0:
                full_ans = ans[(scores - full).abs() < 1e-9]
                full_ans = full_ans[full_ans != ""]
                uniq = set(full_ans.tolist())
                if len(uniq) == 1:
                    correct = next(iter(uniq))
            multi = bool(ans is not None and any(len(a) > 1 for a in ans.tolist()))
            qtype = "多选题" if (len(correct) > 1 or (not correct and multi)) else "单选题"
            dists = [0.0, 0.0, 0.0, 0.0, 0.0, 0.0]
            if ans is not None:
                n = max(len(ans), 1)
                bucket = {"A": 0, "B": 0, "C": 0, "D": 0, "多选": 0, "未选": 0}
                for a in ans.tolist():
                    if a == "":
                        bucket["未选"] += 1
                    elif len(a) == 1 and a in "ABCD":
                        bucket[a] += 1
                    else:
                        bucket["多选"] += 1
                dists = [bucket[k] / n for k in ("A", "B", "C", "D", "多选", "未选")]
            full_value = int(full) if float(full).is_integer() else full
            rows.append([f"第{label}题", qtype, full_value if full > 0 else 0,
                         rate, mean, correct, *dists])
        sheets.append(ClassSheet(sheet_name, header_row, subheader_row, rows))
    return sheets

def run(
    config_path: str = "config/config.yaml",
    plugin_config: str | None = None,
    input_dir: str | None = None,
    output_dir: str | None = None,
    baseline: str | None = None,
) -> int:
    """执行客观题得分明细汇总与距平分析。"""
    raw = _load_plugin_config(plugin_config)
    config = load_config(config_path)

    semester = raw.get("semester") or config.current_semester
    if not semester:
        raise ValueError("未指定学期（插件配置 semester 或全局 current_semester）")
    low_score_flag = float(raw.get("low_score_flag", 0.6))
    if not (0 < low_score_flag <= 1):
        raise ValueError(f"low_score_flag 应在 (0,1] 之间，当前为 {low_score_flag}")
    try:
        max_errors = int(raw.get("max_date_input_errors", 2))
    except (TypeError, ValueError):
        raise ValueError(
            f"max_date_input_errors 应为整数，当前为 {raw.get('max_date_input_errors')!r}"
        )
    if max_errors < 1:
        raise ValueError(f"max_date_input_errors 应 >= 1，当前为 {max_errors}")

    try:
        src, mode = _resolve_input(input_dir, raw, max_errors)
    except InputDirAborted:
        print("[客观题] 输入目录错误次数过多，已中止，未生成任何文件")
        return 1

    if mode == "class_files":
        prefix, _ = _split_folder_name(src.name)
        exam_name, subject = derive_exam_info(
            src.name, semester, config, raw.get("subject")
        )
        try:
            iso_date = _resolve_exam_date(raw, config, exam_name, semester, subject, prefix)
        except ExamDateAborted:
            print("[客观题] 未确认使用今日日期，已中止，未生成任何文件")
            return 1
        exam_date = _format_exam_date(iso_date)
        out = Path(_resolve_output_dir(output_dir, raw, exam_name))
        out.mkdir(parents=True, exist_ok=True)
        print(f"[客观题] 输入目录: {src}")
        print(f"[客观题] 考试规范名称: {exam_name}（学科: {subject or '未识别'}）")
        files = discover_class_files(src)
        sheets: list[ClassSheet] = []
        for path in files:
            print(f"[客观题] 处理文件: {path}")
            sheets.append(read_class_sheet(path))
        if not any(s.class_name == CLASS_SUM_NAME for s in sheets):
            raise ValueError(f"输入文件夹缺少 {CLASS_SUM_NAME}.xls")
    else:
        subject = raw.get("subject") or detect_subject_from_filename(
            src.name, config.subjects, config.subject_aliases
        )
        if not subject:
            raise ValueError(
                "无法识别科目，请在插件配置中填写 subject 或在原始文件名中包含科目名"
            )
        entry = _match_detail_exam(config, semester, subject, src.name)
        exam_name = entry.name if entry is not None else _detail_fallback_name(
            src, semester, subject, config
        )
        if not exam_name:
            raise ValueError(
                "无法推导考试名称：请在 config/exams 中登记对应考试条目，"
                "或在原始文件名中包含考试名称"
            )
        if raw.get("exam_date"):
            iso_date = _parse_date(raw["exam_date"])
        elif entry is not None and getattr(entry, "date", None):
            iso_date = entry.date
        else:
            try:
                iso_date = _resolve_exam_date(
                    raw, config, exam_name, semester, subject, src.stem
                )
            except ExamDateAborted:
                print("[客观题] 未确认使用今日日期，已中止，未生成任何文件")
                return 1
        exam_date = _format_exam_date(iso_date)
        out = Path(_resolve_output_dir(output_dir, raw, exam_name))
        out.mkdir(parents=True, exist_ok=True)
        print(f"[客观题] 小题分文件: {src}")
        print(f"[客观题] 考试规范名称: {exam_name}（学科: {subject}）")
        school_override = (raw.get("school") or "").strip() or None
        try:
            if school_override:
                _validate_detail_school(src, school_override)
                filter_school = school_override
            else:
                filter_school = _resolve_detail_school(src, config, max_errors)
        except SchoolInputAborted:
            print("[客观题] 学校选择错误次数过多，已中止，未生成任何文件")
            return 1
        sheets = build_detail_class_sheets(
            src, config, semester, subject, school=filter_school
        )
        if not any(s.class_name == CLASS_SUM_NAME for s in sheets):
            raise ValueError("内部错误：小题分流程缺少 全部班级 sheet")

    class_names = [s.class_name for s in sheets if s.class_name != CLASS_SUM_NAME]
    summary_types = _parse_group_types(raw.get("summary_groups"), "summary_groups")
    deviation_types = _parse_group_types(
        raw.get("deviation_groups"), "deviation_groups"
    )
    if raw.get("deviation_groups") is None:
        deviation_types = ["teacher"]  # 默认按任课教师生成距平文件
    # 始终构建教师/层次两类分组：分组类型只决定生成哪些列/文件，基线可引用任意分组
    groups_by_type = build_groups(
        config, semester, subject, class_names, ["teacher", "level"]
    )

    deviation_baseline = resolve_baseline(
        baseline
        or raw.get("baseline")
        or raw.get("bl")
        or CLASS_SUM_NAME,
        groups_by_type,
    )

    for t in summary_types:
        if not groups_by_type.get(t):
            print(
                f"[客观题] 警告: 汇总分组类型 {t} 无可用分组"
                "（请检查 config/subjects / config/classes 配置）"
            )
    for t in deviation_types:
        if not groups_by_type.get(t):
            print(
                f"[客观题] 警告: 距平分组类型 {t} 无可用分组"
                "（请检查 config/subjects / config/classes 配置）"
            )
    summary_units = _build_summary_units(sheets, groups_by_type, summary_types)
    deviation_group_names = [
        g.name for t in deviation_types for g in groups_by_type.get(t, [])
    ]
    print(
        "[客观题] 汇总分组: "
        f"{summary_types or ['无']}"
        f"（文件: {'、'.join(label for label, _ in summary_units) or '无'}）"
    )
    print(
        "[客观题] 距平分组: "
        f"{deviation_types or ['无']}（文件: {'、'.join(deviation_group_names) or '无'}）"
    )
    print(
        "[客观题] 基线: "
        f"{CLASS_SUM_NAME if deviation_baseline.group_type == 'all' else deviation_baseline.name}"
    )

    for label, unit_sheets in summary_units:
        sum_path = out / f"{exam_name}_客观题得分汇总（{label}）.xlsx"
        build_summary_workbook(unit_sheets, exam_name, exam_date).save(sum_path)
        print(f"[客观题] 已保存: {sum_path}")

    deviation_groups: list[GroupDef] = []
    for t in deviation_types:
        deviation_groups.extend(groups_by_type.get(t, []))
    if not deviation_groups:
        print("[客观题] 未配置可用的距平分组（deviation_groups），跳过距平分析")
    else:
        score_df = _build_score_frame(sheets)
        for group in deviation_groups:
            dev_path = out / f"{exam_name}_客观题得分率距平（{group.name}）.xlsx"
            build_deviation_workbook(
                score_df,
                group,
                deviation_baseline,
                low_score_flag,
                exam_name,
                exam_date,
            ).save(dev_path)
            print(f"[客观题] 已保存: {dev_path}")
    return 0
