"""个人成绩单生成。

每个学生一个"成绩条"（表头+数据 2 行，条间空行）；
范围由 config/results/config.yaml 的 personal.scope 决定；
样式按参考文件（Calibri 11、行高 20、参考列宽、纵向打印）。
"""

from __future__ import annotations

from pathlib import Path

import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .config import AnalysisConfig, ExamConfig
from .consolidate import date_range_suffix
from .report import _display_qid, _subjective_pivot
from .result_config import ResultsConfig

_PAPER_SIZE = {"A4": 9, "A3": 8, "Letter": 1}
# 纸张尺寸（英寸）：用于分页容量估算
_PAPER_DIMS = {
    "A4": {"width": 8.27, "height": 11.69},
    "A3": {"width": 11.69, "height": 16.54},
    "Letter": {"width": 8.5, "height": 11.0},
}
def _type_score_cols(valid: pd.DataFrame) -> list[str]:
    """得分表动态题型分列显示名（去掉 分 后缀，排除 客观分/主观分/满分列）。"""
    return [
        c[:-1]
        for c in valid.columns
        if c.endswith("分")
        and not c.endswith("满分")
        and c not in ("客观分", "主观分")
    ]


def _teacher_sets(config: AnalysisConfig, exam: ExamConfig) -> dict[str, set[str]]:
    """当前学期+科目范围内，教师名 -> 任教班级集合。"""
    tm = config.teacher_maps.get((exam.semester, exam.subject))
    by_teacher: dict[str, set[str]] = {}
    if not tm:
        return by_teacher
    for cls, code in tm.class_teachers.items():
        name = tm.teacher_names.get(code, code)
        by_teacher.setdefault(name, set()).add(cls)
    return by_teacher


def _label_for(
    classes: set[str],
    all_classes: set[str],
    teacher_sets: dict[str, set[str]],
) -> str:
    if classes == all_classes:
        return "全部班级"
    for teacher, tset in teacher_sets.items():
        if classes == tset:
            return teacher
    return "自定义"


def _resolve_teacher(entry: str, teacher_names: dict[str, str]) -> str:
    """教师代号（A/B/C）转教师名称；已是名称则原样返回；无效则报错。"""
    if entry in teacher_names:
        return teacher_names[entry]
    if entry in set(teacher_names.values()):
        return entry
    raise ValueError(
        f"教师配置无效: {entry!r}（应为教师代号 A/B/C... 或教师名称）"
    )


def _build_strip_rows(
    valid: pd.DataFrame,
    exam: ExamConfig,
    wide: pd.DataFrame,
    big: pd.DataFrame,
    subj_cols: list[str],
    big_cols: list[str],
    classes: set[str],
    cfg: ResultsConfig,
) -> tuple[list[str], list[list[object]], list[float], set[int], dict[int, int]]:
    """构建 表头+数据（每学生一组，含空行），返回 (列名, 行列表, 列宽, 表头行号, 块大小)。"""
    p = cfg.personal
    wide_s = wide.copy()
    wide_s.index = wide_s.index.astype(str)
    big_s = big.copy()
    big_s.index = big_s.index.astype(str)
    subj_by_base: dict[str, list[str]] = {}
    for c in subj_cols:
        # 分组题号（17(1)(2) 等）没有短横线，统一取前导大题号归组
        lead = re.match(r"\d+", c)
        subj_by_base.setdefault(lead.group() if lead else c, []).append(c)
    sub = valid[valid["class_name"].isin(classes)].copy()
    sub = sub.sort_values(
        ["class_name", "total_score", "student_id"],
        ascending=[True, not p.sort_by_score_desc, True],
    )

    type_cols = _type_score_cols(valid)
    cols = ["班级", "姓名", "考试", "班次", "校次", "总分", "客观分", "主观分"] + type_cols
    cw = p.layout.column_widths
    widths = list(cw["first10"][:8]) + [6] * len(type_cols)
    if exam.question_display == "merged":
        merged_cols = [f"{p.merged_prefix}{b}{p.merged_suffix}" for b in big_cols]
        cols += merged_cols
        widths += [cw["merged_question_cols"]] * len(merged_cols)
    else:
        sub_disp = [_display_qid(c) for c in subj_cols]
        cols += sub_disp
        widths += [cw["split_question_cols"]] * len(sub_disp)
    if exam.show_big_questions:
        big_disp = [f"{p.big_score_prefix}{b}{p.big_score_suffix}" for b in big_cols]
        cols += big_disp
        widths += [cw["big_question_cols"]] * len(big_disp)

    blank = p.layout.blank_rows_between
    n = len(sub)
    size_full = 2 + blank
    size_no_blank = 2
    drops = _plan_block_drops(
        [size_full] * n, [size_no_blank] * n, _rows_per_page(cfg)
    )

    rows: list[list[object]] = []
    header_rows: set[int] = set()
    block_sizes: dict[int, int] = {}
    for idx, (_, srow) in enumerate(sub.iterrows()):
        header_row = len(rows) + 1
        header_rows.add(header_row)
        rows.append(list(cols))  # 表头
        sid = str(srow["student_id"])
        data = [
            srow["class_name"],
            srow["name"] if not pd.isna(srow["name"]) else "",
            exam.effective_short_name,
            srow["班次"],
            srow["校次"],
            srow["total_score"],
            srow["objective_score"],
            srow["subjective_score"],
        ]
        for t in type_cols:
            data.append(srow[f"{t}分"])
        if exam.question_display == "merged":
            for base, subs in subj_by_base.items():
                vals = [
                    wide_s.loc[sid, c] if sid in wide_s.index else None
                    for c in subs
                ]
                data.append(
                    "|".join("" if pd.isna(v) else str(int(v)) for v in vals)
                )
        else:
            for c in subj_cols:
                data.append(wide_s.loc[sid, c] if sid in wide_s.index else None)
        if exam.show_big_questions:
            for c in big_cols:
                data.append(big_s.loc[sid, c] if sid in big_s.index else None)
        rows.append(data)
        if not drops[idx]:
            for _ in range(blank):
                rows.append([None] * len(cols))
        # 块大小 = 表头 1 行 + 数据行 + 空行（+1 补上表头行本身）
        block_sizes[header_row] = len(rows) - header_row + 1
    return cols, rows, widths, header_rows, block_sizes


def _write_strip_file(
    exam: ExamConfig,
    valid: pd.DataFrame,
    wide: pd.DataFrame,
    big: pd.DataFrame,
    subj_cols: list[str],
    big_cols: list[str],
    classes: set[str],
    label: str,
    cfg: ResultsConfig,
    config: AnalysisConfig,
) -> str:
    p = cfg.personal
    out_dir = Path(config.results_dir) / (exam.semester or "")
    out_dir.mkdir(parents=True, exist_ok=True)
    date_part = str(exam.date or "").replace("-", "")
    path = out_dir / f"{date_part}_{exam.name}_{label}_个人成绩单.xlsx"

    wb = Workbook()
    wb.remove(wb.active)
    if p.single_sheet:
        ws = wb.create_sheet("个人成绩单")
        cols, rows, widths, header_rows, block_sizes = _build_strip_rows(
            valid, exam, wide, big, subj_cols, big_cols, classes, cfg
        )
        _write_sheet(
            ws, cols, rows, widths, cfg,
            header_rows=header_rows, block_sizes=block_sizes,
        )
    else:
        for cls in sorted(classes):
            sub_valid = valid[valid["class_name"] == cls]
            if sub_valid.empty:
                continue
            ws = wb.create_sheet(str(cls))
            cols, rows, widths, header_rows, block_sizes = _build_strip_rows(
                sub_valid, exam, wide, big, subj_cols, big_cols, {cls}, cfg
            )
            _write_sheet(
                ws, cols, rows, widths, cfg,
                header_rows=header_rows, block_sizes=block_sizes,
            )
    wb.save(path)
    return str(path)


def _write_sheet(
    ws,
    cols: list[str],
    rows: list[list[object]],
    widths: list[float],
    cfg: ResultsConfig,
    header_rows: set[int] | None = None,
    block_sizes: dict[int, int] | None = None,
) -> None:
    lay = cfg.personal.layout
    for row in rows:
        ws.append(row)
    header_font = Font(
        name=lay.font["header"].name,
        size=lay.font["header"].size,
        bold=lay.font["header"].bold,
    )
    data_font = Font(
        name=lay.font["data"].name,
        size=lay.font["data"].size,
        bold=lay.font["data"].bold,
    )
    bold_total = cfg.personal.bold_total_score
    bold_font = Font(
        name=lay.font["data"].name,
        size=lay.font["data"].size,
        bold=True,
    )
    right_from_idx = len(cols)
    if lay.alignment["right_from"] in cols:
        right_from_idx = cols.index(lay.alignment["right_from"])
    center_cols = set(lay.alignment["center_cols"])
    block = 2 + lay.blank_rows_between
    for r in range(1, ws.max_row + 1):
        if header_rows is None:
            pos = (r - 1) % block
            is_header = pos == 0
            is_data = pos == 1
        else:
            is_header = r in header_rows
            is_data = not is_header and any(c.value is not None for c in ws[r])
        for cidx, cell in enumerate(ws[r], start=1):
            if is_header:
                cell.font = header_font
            elif is_data and bold_total and cidx - 1 == 5:
                cell.font = bold_font  # 总分值加粗
            else:
                cell.font = data_font
            if is_header:
                if lay.alignment["header"] == "center_center":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            elif is_data:
                col_name = cols[cidx - 1]
                if col_name in center_cols:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif cidx - 1 >= right_from_idx:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif lay.alignment["data_vertical"] == "center":
                    cell.alignment = Alignment(vertical="center")
        if lay.borders["enabled"] and (is_header or is_data):
            top_style = (
                lay.borders["header_top_style"] if is_header else lay.borders["style"]
            )
            for cell in ws[r]:
                cell.border = Border(
                    top=Side(style=top_style),
                    bottom=Side(style=lay.borders["style"]),
                )
        ws.row_dimensions[r].height = lay.row_height

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    hf = cfg.personal.header_footer
    if hf.enabled:
        ws.oddHeader.left.text = hf.header.left
        ws.oddHeader.center.text = hf.header.center
        ws.oddHeader.right.text = hf.header.right
        ws.oddFooter.left.text = hf.footer.left
        ws.oddFooter.center.text = hf.footer.center
        ws.oddFooter.right.text = hf.footer.right
        hl = hf.fonts["header_left"]
        ws.oddHeader.left.font = hl.name or None
        ws.oddHeader.left.size = hl.size
        hc = hf.fonts["header_center"]
        ws.oddHeader.center.font = hc.name or None
        ws.oddHeader.center.size = hc.size
        hr = hf.fonts["header_right"]
        ws.oddHeader.right.font = hr.name or None
        ws.oddHeader.right.size = hr.size
        fl = hf.fonts["footer_left"]
        ws.oddFooter.left.font = fl.name or None
        ws.oddFooter.left.size = fl.size

    pr = cfg.personal.print
    ws.page_setup.orientation = pr.orientation
    ws.page_setup.paperSize = _PAPER_SIZE.get(pr.paper_size, 9)
    ws.page_margins.left = pr.margin["left"]
    ws.page_margins.right = pr.margin["right"]
    ws.page_margins.top = pr.margin["top"]
    ws.page_margins.bottom = pr.margin["bottom"]
    if pr.fit_to_width:
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # 分页分析：保证同一学生的表头与数据在同一页
    if header_rows is None:
        block = 2 + lay.blank_rows_between
        header_rows = set(range(1, ws.max_row + 1, block))
        block_sizes = {hr: block for hr in header_rows}
    _apply_page_breaks(ws, header_rows, block_sizes, cfg)


def _rows_per_page(cfg: ResultsConfig) -> int:
    """按纸张/边距/行高估算每页可容纳行数。"""
    pr = cfg.personal.print
    if pr.rows_per_page:
        return max(int(pr.rows_per_page), 1)  # 配置优先
    dims = _PAPER_DIMS.get(pr.paper_size, _PAPER_DIMS["A4"])
    usable_inch = (
        dims["width"] if pr.orientation == "landscape" else dims["height"]
    ) - pr.margin["top"] - pr.margin["bottom"]
    usable_pt = max(usable_inch * 72, 1)
    row_h = cfg.personal.layout.row_height or 20
    # 保守余量：估算容量再减 1 行，确保分页符早于 Excel 自动分页边界
    return max(int(usable_pt // row_h) - 1, 1)


def _plan_block_drops(
    sizes_full: list[int],
    sizes_no_blank: list[int],
    capacity: int,
) -> list[bool]:
    """规划每个学生块是否删除末尾空行。

    若某块"不含空行刚好等于每页最大行数"（含空行则超页），
    则删除该块空行留在本页，下一块自然进入下一页。
    """
    drops: list[bool] = []
    used = 0
    for full, no_blank in zip(sizes_full, sizes_no_blank):
        if used + full > capacity and used + no_blank == capacity:
            drops.append(True)
            used += no_blank
        else:
            drops.append(False)
            used = full if used + full > capacity else used + full
    return drops


def _apply_page_breaks(
    ws,
    header_rows: set[int],
    block_sizes: dict[int, int],
    cfg: ResultsConfig,
) -> None:
    """在每个学生块表头之前插入水平分页符，保证表头与数据同页。"""
    from openpyxl.worksheet.pagebreak import Break, RowBreak

    capacity = _rows_per_page(cfg)
    default_size = 2 + cfg.personal.layout.blank_rows_between
    breaks = RowBreak()
    used = 0
    warned = False
    for hr in sorted(header_rows):
        size = block_sizes.get(hr, default_size)
        if size > capacity and not warned:
            print(
                f"[提示] 学生块（{size} 行）超过单页容量（{capacity} 行），"
                f"无法保证表头与数据同页"
            )
            warned = True
        if used + size > capacity and used > 0:
            # Break.id 为 0-based 行索引：表头行 hr（1-based）前分页
            breaks.append(Break(id=hr - 1))
            used = size
        else:
            used += size
    if breaks.brk:
        ws.row_breaks = breaks


def build_personal_strips(
    exam: ExamConfig,
    score: pd.DataFrame,
    questions: pd.DataFrame,
    cfg: ResultsConfig,
    config: AnalysisConfig,
) -> list[str]:
    """按范围生成个人成绩单文件，返回路径列表。"""
    valid = score[score["total_score"].notna()]
    if valid.empty:
        return []
    wide, big, subj_cols, big_cols = _subjective_pivot(questions)
    scope = cfg.personal.scope
    all_classes = set(valid["class_name"])
    teacher_sets = _teacher_sets(config, exam)

    tasks: list[tuple[str, set[str]]] = []
    if scope.mode == "teacher":
        tm = config.teacher_maps.get((exam.semester, exam.subject))
        teacher_names = tm.teacher_names if tm else {}
        for tname in scope.teachers:
            name = _resolve_teacher(tname, teacher_names)
            classes = teacher_sets.get(name, set()) & all_classes
            if classes:
                tasks.append((name, classes))
    elif scope.mode == "custom":
        classes = set(scope.classes) & all_classes
        if classes:
            tasks.append((_label_for(classes, all_classes, teacher_sets), classes))
    else:
        tasks.append(("全部班级", all_classes))

    paths = []
    for label, classes in tasks:
        paths.append(
            _write_strip_file(
                exam, valid, wide, big, subj_cols, big_cols,
                classes, label, cfg, config,
            )
        )
    return paths


def build_merged_personal_strips(
    exams: list[ExamConfig],
    scores: list[pd.DataFrame],
    questions_list: list[pd.DataFrame],
    cfg: ResultsConfig,
    config: AnalysisConfig,
) -> list[str]:
    """多场个人成绩单合并：每个学生一个表头，每场考试一行。

    列固定：班级/姓名/考试/班次/校次/总分/客观分/主观分/单选/多选/主观题；
    主观题列为该场各大题得分合并字符串（竖线分隔）。仅多场（>=2）时生成。
    """
    if not exams or len(exams) < 2:
        return []
    valid_list = [s[s["total_score"].notna()].copy() for s in scores]
    all_classes: set[str] = set()
    for v in valid_list:
        all_classes |= set(v["class_name"])
    teacher_sets = _teacher_sets(config, exams[0])

    tasks: list[tuple[str, set[str]]] = []
    scope = cfg.personal.scope
    if scope.mode == "teacher":
        tm = config.teacher_maps.get((exams[0].semester, exams[0].subject))
        teacher_names = tm.teacher_names if tm else {}
        for tname in scope.teachers:
            name = _resolve_teacher(tname, teacher_names)
            classes = teacher_sets.get(name, set()) & all_classes
            if classes:
                tasks.append((name, classes))
    elif scope.mode == "custom":
        classes = set(scope.classes) & all_classes
        if classes:
            tasks.append((_label_for(classes, all_classes, teacher_sets), classes))
    else:
        tasks.append(("全部班级", all_classes))

    paths: list[str] = []
    for label, classes in tasks:
        paths.append(
            _write_merged_strip_file(
                exams, valid_list, questions_list, classes, label, cfg, config
            )
        )
    return paths


def _write_merged_strip_file(
    exams: list[ExamConfig],
    valid_list: list[pd.DataFrame],
    questions_list: list[pd.DataFrame],
    classes: set[str],
    label: str,
    cfg: ResultsConfig,
    config: AnalysisConfig,
) -> str:
    p = cfg.personal
    out_dir = Path(config.results_dir) / (exams[0].semester or "")
    out_dir.mkdir(parents=True, exist_ok=True)
    suffix = date_range_suffix(exams)
    path = out_dir / f"{suffix}_{label}_个人成绩单.xlsx"

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("个人成绩单")
    cols, rows, widths, header_rows, block_sizes = _build_merged_rows(
        exams, valid_list, questions_list, classes, cfg
    )
    _write_sheet(
        ws, cols, rows, widths, cfg,
        header_rows=header_rows, block_sizes=block_sizes,
    )
    wb.save(path)
    return str(path)


def _build_merged_rows(
    exams: list[ExamConfig],
    valid_list: list[pd.DataFrame],
    questions_list: list[pd.DataFrame],
    classes: set[str],
    cfg: ResultsConfig,
) -> tuple[list[str], list[list[object]], list[float], set[int], dict[int, int]]:
    """构建合并成绩单：每个学生一个表头 + 各场一行数据。

    排序：场次数量多的在前（完整场次 -> 少一场 -> ...），
    组内按 班级/平均总分/姓名/考号；不同场数分组间连续，不额外分页。
    """
    p = cfg.personal

    per_exam: list[tuple[ExamConfig, pd.DataFrame, dict[str, str]]] = []
    for exam, valid, questions in zip(exams, valid_list, questions_list):
        _wide, big, _subj_cols, big_cols = _subjective_pivot(questions)
        big_s = big.copy()
        big_s.index = big_s.index.astype(str)
        subj_str: dict[str, str] = {}
        for sid in big_s.index:
            vals = [
                big_s.loc[sid, c] if sid in big_s.index else None
                for c in big_cols
            ]
            subj_str[sid] = "|".join(
                "" if pd.isna(v) else str(int(v)) for v in vals
            )
        v = valid[valid["class_name"].isin(classes)].copy()
        v["_sid"] = v["student_id"].astype(str)
        per_exam.append((exam, v, subj_str))

    type_cols: list[str] = []
    for _, v, _ in per_exam:
        for t in _type_score_cols(v):
            if t not in type_cols:
                type_cols.append(t)
    cols = (
        ["班级", "姓名", "考试", "班次", "校次", "总分", "客观分", "主观分"]
        + type_cols
        + ["主观题"]
    )
    widths = list(p.layout.column_widths["first10"][:8]) + [6] * len(type_cols) + [
        p.layout.column_widths["merged_question_cols"]
    ]

    all_sids: set[str] = set()
    meta: dict[str, tuple[str, str]] = {}
    total_sum: dict[str, float] = {}
    total_cnt: dict[str, int] = {}
    for _, v, _ in per_exam:
        for _, row in v.iterrows():
            sid = row["_sid"]
            all_sids.add(sid)
            if sid not in meta:
                meta[sid] = (
                    str(row["class_name"]),
                    "" if pd.isna(row["name"]) else str(row["name"]),
                )
            total_sum[sid] = total_sum.get(sid, 0.0) + float(row["total_score"])
            total_cnt[sid] = total_cnt.get(sid, 0) + 1

    def avg(sid: str) -> float:
        return total_sum.get(sid, 0.0) / max(total_cnt.get(sid, 1), 1)

    n_exam_count = {
        s: sum(1 for _, v, _ in per_exam if (v["_sid"] == s).any())
        for s in all_sids
    }
    students = sorted(
        all_sids,
        key=lambda s: (
            -n_exam_count[s],  # 场次数量降序：完整场次在前
            meta[s][0],        # 班级升序
            -avg(s),           # 平均总分降序
            meta[s][1],        # 姓名升序
            s,                 # 考号升序
        ),
    )

    blank = p.layout.blank_rows_between
    n_exams = [n_exam_count[sid] for sid in students]
    sizes_full = [1 + k + blank for k in n_exams]
    sizes_no_blank = [1 + k for k in n_exams]
    drops = _plan_block_drops(sizes_full, sizes_no_blank, _rows_per_page(cfg))

    rows: list[list[object]] = []
    header_rows: set[int] = set()
    block_sizes: dict[int, int] = {}
    for idx, sid in enumerate(students):
        header_row = len(rows) + 1
        header_rows.add(header_row)
        rows.append(list(cols))
        cls, name = meta[sid]
        for exam, v, subj_str in per_exam:
            row = v[v["_sid"] == sid]
            if row.empty:
                continue  # 该生未参加本场
            srow = row.iloc[0]
            row_data = [
                cls,
                name,
                exam.effective_short_name,
                srow["班次"],
                srow["校次"],
                srow["total_score"],
                srow["objective_score"],
                srow["subjective_score"],
            ]
            for t in type_cols:
                key = f"{t}分"
                row_data.append(srow[key] if key in srow.index else None)
            row_data.append(subj_str.get(sid, ""))
            rows.append(row_data)
        if not drops[idx]:
            for _ in range(blank):
                rows.append([None] * len(cols))
        # 块大小 = 表头 1 行 + 数据行 + 空行（+1 补上表头行本身）
        block_sizes[header_row] = len(rows) - header_row + 1
    return cols, rows, widths, header_rows, block_sizes
