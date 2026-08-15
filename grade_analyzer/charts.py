"""图表生成（Excel 内嵌，openpyxl 原生图表）。

四种图表：
- 分数段分布柱状图（人数）
- 多场趋势折线图（平均得分率/及格率/优秀率）
- 班级对比条形图（平均得分率）
- 教师对比条形图（平均得分率）
"""

from __future__ import annotations

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference


def add_distribution_chart(ws, dist_df: pd.DataFrame, anchor: str = "I2") -> None:
    """分数段分布柱状图（人数）。"""
    if dist_df.empty:
        return
    chart = BarChart()
    chart.type = "col"
    chart.title = "分数段分布（人数）"
    chart.y_axis.title = "人数"
    chart.x_axis.title = "分数段"
    n = len(dist_df)
    data = Reference(
        ws,
        min_col=dist_df.columns.get_loc("人数") + 1,
        min_row=1,
        max_row=n + 1,
    )
    cats = Reference(
        ws,
        min_col=dist_df.columns.get_loc("分数段") + 1,
        min_row=2,
        max_row=n + 1,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def add_comparison_chart(
    ws, comp_df: pd.DataFrame, label_col: str, anchor: str = "K2"
) -> None:
    """分组对比条形图（平均得分率），label_col 为类别列（班级/教师）。"""
    if comp_df.empty:
        return
    chart = BarChart()
    chart.type = "bar"
    chart.title = f"{label_col}平均得分率对比"
    chart.y_axis.title = label_col
    chart.x_axis.title = "平均得分率"
    n = len(comp_df)
    data = Reference(
        ws,
        min_col=comp_df.columns.get_loc("平均得分率") + 1,
        min_row=1,
        max_row=n + 1,
    )
    cats = Reference(
        ws,
        min_col=comp_df.columns.get_loc(label_col) + 1,
        min_row=2,
        max_row=n + 1,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def add_trend_chart(ws, trends_df: pd.DataFrame) -> None:
    """多场趋势折线图（平均得分率/及格率/优秀率），置于数据下方。"""
    if trends_df.empty:
        return
    chart = LineChart()
    chart.title = "多场趋势（平均得分率/及格率/优秀率）"
    n = len(trends_df)
    data = Reference(
        ws,
        min_col=trends_df.columns.get_loc("平均得分率") + 1,
        min_row=1,
        max_col=trends_df.columns.get_loc("优秀率") + 1,
        max_row=n + 1,
    )
    cats = Reference(
        ws,
        min_col=trends_df.columns.get_loc("考试名称") + 1,
        min_row=2,
        max_row=n + 1,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{n + 3}")
